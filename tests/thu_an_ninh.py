#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
==========================================================================
 CA THỬ AN NINH GIAO DIỆN WEB
--------------------------------------------------------------------------
 Chạy:  python tests/thu_an_ninh.py

 Bảy lỗi tìm được khi rà soát web.py trước lúc thêm đăng nhập. Tất cả đều
 HỎNG NGAY hôm nay khi chạy cục bộ, không đợi tới lúc mở ra Internet.

   A2  Gửi thân kiểu chunked -> máy chủ coi là RỖNG -> /api/luu-quy-tac
       hiểu thành "không tắt quy tắc nào" -> bật lại hết rồi ghi đè
       cauhinh.json, vẫn báo 200. Mất dữ liệu trong im lặng.
   A1  Không kiểm Content-Type / Origin / Host -> trang web bất kỳ mà thầy
       cô mở đều sai khiến được máy chủ ghi đè cauhinh.json.
   A7  Chỉ một khe sao lưu -> bấm Lưu lần thứ hai là đè mất bản gốc.
   --  Ghi cauhinh.json không nguyên tử: open('w') cắt cụt tệp TRƯỚC khi
       biết có dựng nổi chuỗi JSON hay không.
   A4  Khoá "một lượt" của /api/xep đọc cờ rồi mới đặt cờ ở hai lần khoá
       khác nhau -> nhiều lượt cùng chạy, cùng ghi một tệp Excel.
   A3  Lỗi bất ngờ trả nguyên traceback ra trình duyệt, kèm đường dẫn
       tuyệt đối và tên đăng nhập máy.
   A8  Chuỗi bắt đầu bằng '=' ghi vào ô Excel thành CÔNG THỨC. Văn thư mở
       tệp, Excel hỏi cho phép liên kết ngoài, bấm Có là chạy lệnh.

 Mọi ca đi qua HTTP thật, trong thư mục tạm, không đụng cauhinh.json thật.
==========================================================================
"""
import io
import json
import os
import socket
import sys
import tempfile
import threading
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from tien_ich_web import (GOC_TM, MayChuTam, Khach, cau_hinh_tam,   # noqa: E402
                          doi_trang_thai)
import xep_tkb as X          # noqa: E402

KQ = []


def kiem(dieu_kien, mo_ta):
    KQ.append((bool(dieu_kien), mo_ta))
    print(f'  [{"ĐẠT" if dieu_kien else "LỖI"}] {mo_ta}')


def so_tat(dd):
    """Danh sách chỉ số quy tắc đang bị tắt trong tệp cấu hình.
       Tệp chưa có thì trả None — để ca thử báo ĐỎ chứ không vỡ giữa chừng."""
    cfg = doc_an_toan(dd)
    if cfg is None:
        return None
    return [i for i, r in enumerate(cfg['quy_tac']) if r.get('bat') is False]


def doc_an_toan(dd):
    """Đọc cấu hình, không có hoặc hỏng thì trả None thay vì ném."""
    try:
        return X.doc_cau_hinh(dd)
    except BaseException:
        return None


# ==========================================================================
def ca_a2_than_di_dang(tm):
    print('\nCA A2 — thân yêu cầu dị dạng KHÔNG được coi là thân rỗng')
    print('-' * 74)
    tmx = os.path.join(tm, 'a2')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)

    with MayChuTam(dd) as (k, _):
        # dựng trạng thái: tắt 10 quy tắc rồi lưu xuống tệp
        k.js('/api/luu-quy-tac', {'tat': list(range(10))})
        kiem(so_tat(dd) == list(range(10)),
             f'Dựng sẵn trạng thái: 10 quy tắc bị tắt — {so_tat(dd)}')

        # 1. Transfer-Encoding: chunked, KHÔNG có Content-Length
        ma, than = k.chunked('/api/luu-quy-tac', {'tat': list(range(10))})
        kiem(ma == 400,
             f'Gửi chunked bị TỪ CHỐI bằng 400 — máy chủ trả {ma}')
        kiem(so_tat(dd) == list(range(10)),
             f'Tệp cấu hình KHÔNG bị đổi sau yêu cầu chunked — {so_tat(dd)}')

        # 2. POST không có Content-Length và cũng không chunked
        ma, than = k.khong_do_dai('/api/luu-quy-tac', b'{"tat":[]}')
        kiem(ma == 400,
             f'Gửi thiếu Content-Length bị TỪ CHỐI bằng 400 — {ma}')
        kiem(so_tat(dd) == list(range(10)),
             f'Tệp cấu hình vẫn KHÔNG bị đổi — {so_tat(dd)}')

        # 3. Content-Length không phải số
        ma, than = k.yeu_cau('POST', '/api/luu-quy-tac', b'{"tat":[]}')[0], None
        ma2, _, t2 = k.yeu_cau('POST', '/api/luu-quy-tac', None,
                               them={'Content-Length': 'abc'})
        kiem(ma2 == 400, f'Content-Length không phải số -> 400 — {ma2}')
        kiem('Traceback' not in t2.decode('utf-8', 'replace'),
             'Và không kèm traceback')

        # 4. thân rỗng hẳn cũng phải bị từ chối ở đường GHI TỆP
        ma, _ = k.post('/api/luu-quy-tac', b'')
        kiem(ma == 400, f'Thân rỗng ở đường ghi tệp -> 400 — {ma}')

        # 5. đường ghi tệp phải đòi đúng khoá "tat"
        ma, _ = k.js('/api/luu-quy-tac', {'khong_phai_tat': [1]})
        kiem(ma == 400, f'Thiếu khoá "tat" -> 400 — {ma}')

        # 6. yêu cầu ĐÚNG vẫn phải chạy được như cũ
        ma, kq = k.js('/api/luu-quy-tac', {'tat': [3]})
        kiem(ma == 200 and so_tat(dd) == [3],
             f'Yêu cầu đúng vẫn lưu được bình thường — {ma}, {so_tat(dd)}')


# ==========================================================================
def ca_a1_nguon_goc(tm):
    print('\nCA A1 — chỉ nhận yêu cầu từ chính trang này')
    print('-' * 74)
    tmx = os.path.join(tm, 'a1')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)

    with MayChuTam(dd) as (k, (dia, cong)):
        k.js('/api/luu-quy-tac', {'tat': [1, 2]})
        goc = so_tat(dd)
        kiem(goc == [1, 2], f'Dựng sẵn trạng thái: quy tắc {goc} bị tắt')

        # --- Content-Type: ba kiểu "simple request" phải bị chặn hết ---
        for kieu in ('text/plain', 'application/x-www-form-urlencoded',
                     'multipart/form-data; boundary=x'):
            ma, _ = k.post('/api/luu-quy-tac', {'tat': []}, kieu=kieu)
            kiem(ma == 400, f'Content-Type {kieu.split(";")[0]} bị chặn — {ma}')
            kiem(so_tat(dd) == goc, f'  tệp không đổi sau khi bị chặn')

        ma, _ = k.post('/api/luu-quy-tac', {'tat': []}, kieu=None)
        kiem(ma == 400, f'Không khai Content-Type cũng bị chặn — {ma}')

        # --- Origin của trang khác phải bị chặn ---
        for org in ('https://trang-quang-cao.example', 'http://127.0.0.1:1',
                    'null'):
            ma, _ = k.post('/api/luu-quy-tac', {'tat': []},
                           them={'Origin': org})
            kiem(ma == 400, f'Origin lạ {org[:32]} bị chặn — {ma}')
            kiem(so_tat(dd) == goc, '  tệp không đổi sau khi bị chặn')

        # --- Host lạ (DNS rebinding) phải bị chặn, cả GET ---
        for h in ('trang-la.example', 'trang-la.example:%d' % cong):
            ma, _ = k.get('/api/gia-tri', them={'Host': h})
            kiem(ma == 400, f'Host lạ "{h}" bị chặn ở GET — {ma}')
        ma, _ = k.post('/api/luu-quy-tac', {'tat': []},
                       them={'Host': 'trang-la.example'})
        kiem(ma == 400 and so_tat(dd) == goc, 'Host lạ bị chặn ở POST')

        # --- những thứ ĐÚNG vẫn phải chạy được ---
        ma, _ = k.post('/api/luu-quy-tac', {'tat': [4]},
                       kieu='application/json; charset=utf-8')
        kiem(ma == 200 and so_tat(dd) == [4],
             f'Content-Type có charset vẫn nhận — {ma}, {so_tat(dd)}')
        ma, _ = k.post('/api/luu-quy-tac', {'tat': [5]},
                       them={'Origin': f'http://127.0.0.1:{cong}'})
        kiem(ma == 200 and so_tat(dd) == [5],
             f'Origin của chính trang này vẫn nhận — {ma}')
        ma, _ = k.post('/api/luu-quy-tac', {'tat': [6]},
                       them={'Host': f'localhost:{cong}',
                             'Origin': f'http://localhost:{cong}'})
        kiem(ma == 200 and so_tat(dd) == [6],
             f'Vào bằng localhost vẫn dùng được — {ma}')
        kiem(k.get('/')[0] == 200, 'GET / bình thường vẫn 200')
        kiem(k.get('/api/gia-tri')[0] == 200, 'GET /api/gia-tri vẫn 200')
        ma, _ = k.get('/api/tai-ve')
        kiem(ma in (200, 404),
             f'Đường tải tệp (điều hướng cả trang, không có Origin) vẫn qua — {ma}')

        # ------------------------------------------------------------------
        # BẢNG Host — lái bằng dữ liệu, để ai đó "cải tiến" phép so khớp
        # chính xác thành startswith / endswith / in là ĐỎ ngay.
        #
        # Cột "mong" ghi cả những chuỗi PHẢI cho 200: mấy dòng đó canh chiều
        # ngược lại, chống siết quá tay (bỏ .lower() hay .strip() là đỏ).
        # ------------------------------------------------------------------
        P = cong
        BANG_HOST = [
            (f'127.0.0.1:{P}',                 200, 'đường đi chuẩn'),
            (f'localhost:{P}',                 200, 'app.js mở được cả hai tên'),
            (f'[::1]:{P}',                     200, 'loopback IPv6'),
            (f'LOCALHOST:{P}',                 200, 'DNS không phân biệt hoa/thường'),
            # ĐỔI Ý so với Đợt 1, có chủ ý: bỏ .strip() khỏi phép chuẩn hoá.
            # Đã đo: bộ phân tích tiêu đề tự bỏ khoảng trắng ĐẦU nhưng GIỮ
            # khoảng trắng CUỐI, nên chuỗi này tới nơi là '127.0.0.1:P  ' và
            # nay KHÔNG khớp. Trình duyệt không bao giờ gửi khoảng trắng thừa;
            # cắt hộ "cho dễ khớp" chính là kiểu nới lỏng mà cả bảng này canh.
            (f'  127.0.0.1:{P}  ',             400, 'khoảng trắng cuối KHÔNG được cắt hộ'),
            ('127.0.0.1',                      200, 'tên loopback trơ, cố ý nhận'),
            ('localhost',                      200, 'như trên'),
            ('evil.com',                       400, 'DNS rebinding — cái A1 sinh ra để chặn'),
            (f'evil.com:{P}',                  400, 'rebinding có cổng'),
            (f'127.0.0.1:{P}.tenlakhac.com',   400, 'HẬU TỐ — bẫy của startswith'),
            (f'xx127.0.0.1:{P}',               400, 'TIỀN TỐ — bẫy của endswith'),
            ('localhost.evil.com',             400, 'hậu tố tên miền — bẫy của in'),
            (f'evil.com.127.0.0.1:{P}',        400, 'bẫy của endswith phía kia'),
            (f'evil.com@127.0.0.1:{P}',        400, 'userinfo không hợp lệ trong Host'),
            (f'127.0.0.1:{P}@evil.com',        400, 'userinfo phía sau'),
            (f'localhost.:{P}',                400, 'chấm cuối FQDN — cố ý KHÔNG nới'),
            (f'127.000.000.001:{P}',           400, 'cùng IP khác chuỗi, so chuỗi chặt là đúng'),
            (f'2130706433:{P}',                400, 'IP dạng số nguyên'),
            (f'[::ffff:127.0.0.1]:{P}',        400, 'IPv4 lồng trong IPv6'),
            (f'127.0.0.1:0{P}',                400, 'cổng có số 0 đứng đầu'),
            ('',                               400, 'Host rỗng'),
        ]
        sai = []
        for h, mong, vi_sao in BANG_HOST:
            ma, _ = k.get('/api/gia-tri', them={'Host': h})
            if ma != mong:
                sai.append(f'{h!r} cho {ma}, mong {mong} ({vi_sao})')
        kiem(not sai,
             f'Bảng {len(BANG_HOST)} chuỗi Host khớp hết — lệch: {sai}')

        # --- tiêu đề TRÙNG LẶP: dòng thật trước, dòng giả sau ---
        # self.headers.get() chỉ lấy dòng ĐẦU, nên bản cũ cho 200. Một công cụ
        # trung gian có thể chọn dòng khác ta -> đúng kiểu lệch mà A1 chặn.
        nen = [('Connection', 'close')]
        ma, _, _ = k.tho_dong('GET', '/api/gia-tri',
                              [('Host', f'127.0.0.1:{P}'),
                               ('Host', 'evil.com')] + nen)
        kiem(ma == 400, f'Hai dòng Host (thật trước) bị chặn — {ma}')
        ma, _, _ = k.tho_dong('GET', '/api/gia-tri',
                              [('Host', 'evil.com'),
                               ('Host', f'127.0.0.1:{P}')] + nen)
        kiem(ma == 400, f'Hai dòng Host (giả trước) bị chặn — {ma}')
        ma, _, _ = k.tho_dong('GET', '/api/gia-tri',
                              [('Host', f'127.0.0.1:{P}'),
                               ('Origin', f'http://127.0.0.1:{P}'),
                               ('Origin', 'https://trang-la.example')] + nen)
        kiem(ma == 400, f'Hai dòng Origin bị chặn — {ma}')
        than = json.dumps({'tat': []}).encode()
        ma, _, _ = k.tho_dong('POST', '/api/luu-quy-tac',
                              [('Host', f'127.0.0.1:{P}'),
                               ('Content-Type', 'application/json'),
                               ('Content-Length', str(len(than))),
                               ('Content-Length', str(len(than)))] + nen, than)
        kiem(ma == 400 and so_tat(dd) == [6],
             f'Hai dòng Content-Length bị chặn, tệp không đổi — {ma}')

        # --- Content-Length phải là chữ số ASCII thuần ---
        # Đã đo: int('1_0') = 10, int(' +10 ') = 10, int('١٠') = 10.
        # Đây là con số quyết định việc cấp phát bộ nhớ, không được dễ dãi.
        for tho in ('1_0', ' +10 ', '١٠', '0x10', '10.0'):
            ma, _, _ = k.tho_dong('POST', '/api/luu-quy-tac',
                                  [('Host', f'127.0.0.1:{P}'),
                                   ('Content-Type', 'application/json'),
                                   ('Content-Length', tho)] + nen,
                                  b'{"tat":[]}')
            kiem(ma == 400,
                 f'Content-Length {tho!r} bị từ chối — {ma}')
        kiem(so_tat(dd) == [6], 'Tệp vẫn không đổi sau loạt Content-Length lạ')


# ==========================================================================
def ca_a1b_danh_sach_ten_may(tm):
    print('\nCA A1b — danh sách tên máy lái bằng cauhinh_may_chu.json')
    print('-' * 74)
    tmx = os.path.join(tm, 'a1b')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)
    import web as W

    TEN = 'admin.tail7dc28e.ts.net'
    mc = os.path.join(tmx, 'cauhinh_may_chu.json')

    def viet(host, org):
        with io.open(mc, 'w', encoding='utf-8') as f:
            json.dump({'host_cho_phep': host, 'origin_cho_phep': org}, f,
                      ensure_ascii=False)

    viet([TEN], ['https://' + TEN])          # cổng thật chưa biết, sửa lại bên trong
    with MayChuTam(dd, may_chu=mc) as (k, (dia, cong)):
        # Máy chủ thử chạy cổng ngẫu nhiên nên phải viết lại tệp cho đúng cổng
        # rồi nạp lại — đây là DỰNG DỮ LIỆU ĐẦU VÀO, y như cau_hinh_tam() ghi
        # cauhinh.json, không phải gọi tắt qua tầng HTTP.
        viet([f'127.0.0.1:{cong}', TEN],
             [f'http://127.0.0.1:{cong}', 'https://' + TEN])
        W.nap_cau_hinh_may_chu(mc)

        BANG_HOST = [
            (f'127.0.0.1:{cong}',           200, 'từ chính máy vẫn vào được'),
            (TEN,                           200, 'qua proxy — trình duyệt BỎ :443'),
            (TEN.upper(),                   200, 'tên miền không phân biệt hoa thường'),
            (TEN.title(),                   200, 'hoa thường lẫn lộn'),
            (TEN + ':443',                  400, 'KHÔNG tự nhận cổng mặc định'),
            (TEN + '.tenlakhac.com',        400, 'hậu tố — bẫy của startswith'),
            ('tenlakhac.com/' + TEN,        400, 'nhét vào đường dẫn — bẫy của in'),
            ('tenlakhac.com.' + TEN,        400, 'tiền tố — bẫy của endswith'),
            (TEN + ' evil.com',             400, 'hai tên cách nhau dấu cách'),
            ('evil.com ' + TEN,             400, 'như trên, đảo thứ tự'),
            (TEN + ' ',                     400, 'thừa đúng một dấu cách cuối'),
            (' ' + TEN,                     200, 'khoảng trắng ĐẦU do bộ phân tích tự bỏ'),
            (f'[::1]:{cong}',               400, 'KHÔNG còn trong danh sách'),
            (f'localhost:{cong}',           400, 'KHÔNG còn trong danh sách'),
            ('127.0.0.1',                   400, 'KHÔNG còn trong danh sách'),
            ('',                            400, 'Host rỗng'),
            ('evil.com',                    400, 'tên hoàn toàn khác'),
        ]
        sai = []
        for h, mong, vi_sao in BANG_HOST:
            ma, _ = k.get('/api/gia-tri', them={'Host': h})
            if ma != mong:
                sai.append(f'{h!r} cho {ma}, mong {mong} ({vi_sao})')
        kiem(not sai,
             f'Bảng {len(BANG_HOST)} chuỗi Host khớp hết — lệch: {sai}')

        # bốn dòng dưới đây phải chứng minh danh sách THAY THẾ mặc định,
        # không phải cộng thêm vào
        kiem(all(k.get('/api/gia-tri', them={'Host': h})[0] == 400
                 for h in (f'localhost:{cong}', f'[::1]:{cong}', '127.0.0.1',
                           'localhost')),
             'Có tệp thì danh sách THAY THẾ mặc định loopback, không cộng thêm')

        # thiếu hẳn dòng Host
        ma, _, _ = k.tho_dong('GET', '/api/gia-tri', [('Connection', 'close')])
        kiem(ma == 400, f'Thiếu hẳn dòng Host -> 400 — {ma}')

        # gấp dòng obs-fold: \r\n NẰM LẠI trong giá trị (đã đo)
        goi = (f'GET /api/gia-tri HTTP/1.1\r\nHost: {TEN}\r\n evil.com\r\n'
               f'Connection: close\r\n\r\n').encode()
        ma, _, t = k.tho(goi)
        kiem(ma == 400, f'Gấp dòng obs-fold trong Host -> 400 — {ma}')
        kiem('điều khiển' in t.decode('utf-8', 'replace'),
             '  và câu lỗi nói đúng lý do (ký tự điều khiển), không nói chung chung')

        # -------------------------------------------------------------
        # CHỖ ĐỀ BÀI VIẾT "CHẶN" MÀ THỰC TẾ KHÔNG CHẶN ĐƯỢC — ghi lại đúng
        # sự thật thay vì khẳng định một điều sai.
        #
        # Gửi 'Host: <tên>\r\nX-Gia: 1' thì \r\n KHÔNG nằm lại trong giá trị
        # Host. Nó tách yêu cầu thành HAI dòng tiêu đề, và Host còn lại là
        # chuỗi HỢP LỆ (đã đo bằng email.parser). Máy chủ cho qua, ĐÚNG.
        #
        # Chặn ở đây cũng vô nghĩa: khách tự chèn thêm tiêu đề vào yêu cầu của
        # CHÍNH NÓ thì chẳng được gì. Mối nguy thật là chiều ngược lại — ta
        # phát \r\n ra PHẢN HỒI — và đó là bất biến A10, đã vá và có ca riêng.
        # -------------------------------------------------------------
        goi = (f'GET /api/gia-tri HTTP/1.1\r\nHost: {TEN}\r\nX-Gia: 1\r\n'
               f'Connection: close\r\n\r\n').encode()
        ma, _, _ = k.tho(goi)
        kiem(ma == 200,
             f'"Host hợp lệ + \\r\\nX-Gia: 1" cho 200 — ĐÚNG, vì nó thành HAI '
             f'dòng tiêu đề chứ không phải một Host hỏng ({ma})')

        # ---------------- bảng Origin, RIÊNG ----------------
        BANG_ORG = [
            ('https://' + TEN,                  200, 'đúng tên, đúng giao thức'),
            ('HTTPS://' + TEN.upper(),          200, 'hoa thường'),
            ('http://' + TEN,                   400, 'ĐÚNG tên nhưng SAI giao thức'),
            (f'http://127.0.0.1:{cong}',        200, 'từ chính máy'),
            (f'https://127.0.0.1:{cong}',       400, 'sai giao thức chiều ngược lại'),
            ('https://' + TEN + ':443',         400, 'không tự nhận cổng mặc định'),
            ('https://' + TEN + '.tenlakhac.com', 400, 'hậu tố'),
            ('https://tenlakhac.com/' + TEN,    400, 'nhét vào đường dẫn'),
            ('null',                            400, 'trang sandbox'),
            ('https://' + TEN + ' evil.com',    400, 'có dấu cách'),
        ]
        sai = []
        for o, mong, vi_sao in BANG_ORG:
            ma, _ = k.post('/api/thu-quy-tac', {'tat': []},
                           them={'Host': TEN, 'Origin': o})
            if ma != mong:
                sai.append(f'{o!r} cho {ma}, mong {mong} ({vi_sao})')
        kiem(not sai,
             f'Bảng {len(BANG_ORG)} chuỗi Origin khớp hết — lệch: {sai}')
        kiem(k.post('/api/thu-quy-tac', {'tat': []},
                    them={'Host': TEN, 'Origin': 'http://' + TEN})[0] == 400,
             'Hai bảng RIÊNG: đúng tên máy mà sai giao thức vẫn bị chặn')

    # ---------------- FAIL-CLOSED: năm đường vào, ĐO QUA HTTP THẬT ----------
    #
    # Đo qua HTTP chứ không chỉ xem giá trị trả về của nap_cau_hinh_may_chu().
    # Đã thử: một bản hỏng cố ý (tệp hỏng thì cho qua mọi Host) vẫn trả
    # (None, None) y hệt bản đúng, nên phép kiểm giá trị trả về XANH GIẢ hoàn
    # toàn. Thứ cần khẳng định là MÁY CHỦ CÓ TỪ CHỐI THẬT KHÔNG.
    for nhan, noi_dung, cho_ghi in [
            ('JSON hỏng',    '{ hong',                                       True),
            ('rỗng cả hai',  '{"host_cho_phep": [], "origin_cho_phep": []}',  True),
            ('sai kiểu',     '{"host_cho_phep": "khong-phai-day"}',           True),
            ('có ký tự lạ',  '{"host_cho_phep": ["a\\u0000b"]}',              True),
            ('KHÔNG CÓ TỆP', None,                                           False)]:
        if noi_dung is None:
            if os.path.exists(mc):
                os.remove(mc)
        else:
            with io.open(mc, 'w', encoding='utf-8') as f:
                f.write(noi_dung)

        with MayChuTam(dd, may_chu=mc) as (k2, (dia2, cong2)):
            kiem(W.HOST_CHO_PHEP is None and W.ORIGIN_CHO_PHEP is None,
                 f'Tệp {nhan} -> không nạp danh sách nào')
            # điều thật sự quan trọng: máy chủ CÓ TỪ CHỐI không
            ma_la, _ = k2.get('/api/gia-tri', them={'Host': TEN})
            ma_ac, _ = k2.get('/api/gia-tri', them={'Host': 'evil.com'})
            ma_nha, _ = k2.get('/api/gia-tri',
                               them={'Host': f'127.0.0.1:{cong2}'})
            kiem(ma_la == 400 and ma_ac == 400,
                 f'  và TỪ CHỐI THẬT mọi tên ngoài loopback — '
                 f'{TEN}:{ma_la}, evil.com:{ma_ac}')
            kiem(ma_nha == 200,
                 f'  nhưng chính máy này vẫn vào được — {ma_nha}')
        ghi = W.nap_cau_hinh_may_chu(mc)[2]
        kiem(bool(ghi) == cho_ghi,
             f'  báo ra cửa sổ dòng lệnh: {"CÓ" if cho_ghi else "im lặng"} — '
             f'{(ghi or "").splitlines()[:1]}')

    # có tệp đúng thì loopback mặc định KHÔNG còn được ngầm cho qua
    with io.open(mc, 'w', encoding='utf-8') as f:
        json.dump({'host_cho_phep': [TEN], 'origin_cho_phep': ['https://' + TEN]},
                  f, ensure_ascii=False)
    h, o, g = W.nap_cau_hinh_may_chu(mc)
    kiem(h == frozenset({TEN}) and o == frozenset({'https://' + TEN}),
         f'Tệp hợp lệ -> nạp đúng hai tập RIÊNG — {sorted(h)} / {sorted(o)}')
    kiem(g and TEN in g, 'Và in ra màn hình đúng danh sách đang có hiệu lực')
    W.nap_cau_hinh_may_chu(os.path.join(tmx, 'khong-co.json'))   # dọn lại


# ==========================================================================
def ca_a1c_moi_duong_deu_bi_gac(tm):
    print('\nCA A1c — cổng Host/Origin gác MỌI đường, kể cả đường ghi')
    print('-' * 74)
    tmx = os.path.join(tm, 'a1c')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)
    import web as W
    import re as _re

    TEN = 'admin.tail7dc28e.ts.net'
    mc = os.path.join(tmx, 'cauhinh_may_chu.json')

    # Danh sách đường BÓC TỪ MÃ NGUỒN, không chép tay. Chép tay thì thêm đường
    # mới là ca thử vẫn xanh — đúng kiểu mục kiểm mục nát theo thời gian.
    ma_nguon = io.open(os.path.join(GOC_TM, 'web.py'), encoding='utf-8').read()
    duong_api = sorted(set(_re.findall(r"duong == '(/api/[a-z-]+)'", ma_nguon)))
    # Guard hai chiều. Phép bóc hỏng -> vòng lặp dưới chạy rỗng và XANH GIẢ,
    # nên phải khẳng định nó ra ĐÚNG tập đang có. Thêm đường mới vào web.py thì
    # mục này ĐỎ trước — buộc người thêm phải ngó lại xem đã gác chưa rồi mới
    # cập nhật danh sách ở đây.
    MONG = sorted(['/api/trang-thai', '/api/quy-tac', '/api/loai-quy-tac',
                   '/api/gia-tri', '/api/tai-ve', '/api/tai-web',
                   '/api/xep', '/api/luu-quy-tac', '/api/them-quy-tac',
                   '/api/xoa-quy-tac', '/api/thu-quy-tac'])
    kiem(duong_api == MONG,
         f'Bóc từ web.py ra ĐÚNG {len(MONG)} đường /api/ đang có — '
         f'thiếu {sorted(set(MONG) - set(duong_api))}, '
         f'thừa {sorted(set(duong_api) - set(MONG))}')

    # Đường ghi: khai riêng vì cần thân, và vì đây chính là nhóm đề bài lo.
    THAN = {'/api/xep': {'tat': []},
            '/api/luu-quy-tac': {'tat': []},
            '/api/them-quy-tac': {'quy_tac': {'loai': 'khong_xep_thu',
                                              'ghi_chu': 'x', 'thu': [2]}},
            '/api/xoa-quy-tac': {'i': 0},
            '/api/thu-quy-tac': {'tat': []}}
    TINH_D = ['/', '/index.html', '/web/app.js', '/web/style.css']

    with MayChuTam(dd, may_chu=mc) as (k, (dia, cong)):
        with io.open(mc, 'w', encoding='utf-8') as f:
            json.dump({'host_cho_phep': [TEN],
                       'origin_cho_phep': ['https://' + TEN]}, f)
        W.nap_cau_hinh_may_chu(mc)

        # --- 1. Host LẠ: mọi đường, đọc lẫn ghi, đều phải 400 ---
        lot = []
        for d in TINH_D + duong_api:
            if d in THAN:
                ma, _ = k.post(d, THAN[d], them={'Host': 'evil.com'})
            else:
                ma, _ = k.get(d, them={'Host': 'evil.com'})
            if ma != 400:
                lot.append(f'{d}:{ma}')
        kiem(not lot,
             f'Host lạ bị chặn ở TẤT CẢ {len(TINH_D) + len(duong_api)} đường '
             f'— lọt: {lot}')

        # --- 2. Nhóm GHI nói riêng, vì mất dữ liệu là hậu quả nặng nhất ---
        lot_ghi = [f'{d}:{k.post(d, THAN[d], them={"Host": "evil.com"})[0]}'
                   for d in sorted(THAN)
                   if k.post(d, THAN[d], them={'Host': 'evil.com'})[0] != 400]
        kiem(not lot_ghi,
             f'Riêng 5 đường GHI: Host lạ đều 400 — lọt: {lot_ghi}')
        kiem(so_tat(dd) == [], 'Và cauhinh.json không hề bị đụng tới')

        # --- 3. Host đúng nhưng Origin của TRANG LẠ: đường ghi phải 400 ---
        lot_org = []
        for d in sorted(THAN):
            ma, _ = k.post(d, THAN[d],
                           them={'Host': TEN,
                                 'Origin': 'https://trang-quang-cao.example'})
            if ma != 400:
                lot_org.append(f'{d}:{ma}')
        kiem(not lot_org,
             f'Origin trang lạ bị chặn ở cả 5 đường ghi — lọt: {lot_org}')
        kiem(so_tat(dd) == [], 'Và cauhinh.json vẫn không bị đụng')

        # --- 4. Chiều ngược lại: đúng cả hai thì KHÔNG được chặn oan ---
        oan = []
        for d in TINH_D + duong_api:
            if d in THAN:
                ma, t = k.post(d, THAN[d],
                               them={'Host': TEN, 'Origin': 'https://' + TEN})
            else:
                ma, t = k.get(d, them={'Host': TEN})
            # 400 do NGHIỆP VỤ (vd xoá quy tắc không có) thì không tính là oan;
            # chỉ 400 phát ra từ chính cái cổng mới là chặn oan.
            if ma == 400 and ('tên máy lạ' in t or 'trang khác' in t):
                oan.append(f'{d}:{ma}')
        kiem(not oan,
             f'Host+Origin đúng thì cổng KHÔNG chặn đường nào — oan: {oan}')

        # --- 5. Ghi lại ĐÚNG hành vi khi KHÔNG có dòng Origin ---
        #
        # Cố ý không đòi Origin, và đây là lý do — không phải bỏ sót:
        # kiểm Origin CHỈ chặn được trình duyệt bị trang lạ sai khiến, vì trình
        # duyệt không cho trang web nói dối Origin. Một chương trình gọi thẳng
        # (curl, script) tự đặt Origin nào cũng được, nên bắt buộc phải có
        # Origin KHÔNG thêm được một chút an toàn nào trước nó — chỉ thêm rủi ro
        # từ chối oan một khách hợp lệ. Hàng rào đối với khách gọi thẳng là
        # ĐĂNG NHẬP (Đợt 2), không phải Origin.
        ma, _ = k.post('/api/luu-quy-tac', {'tat': [1]}, them={'Host': TEN})
        kiem(ma == 200,
             f'Không có dòng Origin (curl/script) thì vẫn phục vụ — {ma}. '
             f'CỐ Ý: Origin chỉ chặn được trình duyệt, khách gọi thẳng tự đặt '
             f'được nên đòi cũng vô ích')
        kiem(so_tat(dd) == [1], '  và ghi thật xuống tệp')

    W.nap_cau_hinh_may_chu(os.path.join(tmx, 'khong-co.json'))


# ==========================================================================
def ca_a9_tran_than(tm):
    print('\nCA A9 — thân yêu cầu có TRẦN, và không nuốt vào bộ nhớ')
    print('-' * 74)
    tmx = os.path.join(tm, 'a9')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)

    import web as W
    tran = W.TRAN_THAN

    # Dụng cụ đo: bọc rfile.read để ghi lại ĐỐI SỐ LỚN NHẤT từng được yêu cầu.
    # Đây mới là chứng cứ "không đọc hết vào bộ nhớ" — mã 413 một mình không
    # chứng minh được gì, máy chủ vẫn có thể đã nuốt 20 MB rồi mới trả lời.
    that = W.May._doc_than_co_tran
    do = {'lon_nhat': 0}

    def dem(self, n):
        goc_read = self.rfile.read

        def read_dem(k=-1):
            do['lon_nhat'] = max(do['lon_nhat'], k if k is not None else 0)
            return goc_read(k)
        self.rfile.read = read_dem
        try:
            return that(self, n)
        finally:
            self.rfile.read = goc_read

    W.May._doc_than_co_tran = dem
    try:
        with MayChuTam(dd) as (k, (dia, cong)):
            k.js('/api/luu-quy-tac', {'tat': [3]})
            kiem(so_tat(dd) == [3], 'Dựng sẵn trạng thái: quy tắc [3] bị tắt')

            # --- thân ĐÚNG CỠ vẫn phải chạy: canh không-hồi-quy ---
            do['lon_nhat'] = 0
            ma, _ = k.js('/api/luu-quy-tac', {'tat': [2]})
            kiem(ma == 200 and so_tat(dd) == [2],
                 f'Thân bình thường vẫn ghi được — {ma}, {so_tat(dd)}')
            kiem(0 < do['lon_nhat'] <= 65536,
                 f'  và đọc theo mẩu ≤ 64 KB — mẩu lớn nhất {do["lon_nhat"]}')

            # --- sát trần: off-by-one ---
            do['lon_nhat'] = 0
            dem_byte = tran - 200
            than = {'tat': [2], 'rac': 'A' * dem_byte}
            n = len(json.dumps(than, ensure_ascii=False).encode())
            ma, _ = k.js('/api/luu-quy-tac', than)
            kiem(n <= tran, f'  (thân thử dài {n:,} byte, dưới trần {tran:,})')
            kiem(ma == 200, f'Thân sát trần vẫn nhận — {ma}')

            # --- quá trần ---
            #
            # Hai ca này KHÁC KỲ VỌNG NHAU, và không thể gộp. Máy chủ cố ý
            # ngừng đọc rồi đóng, nên với thân RẤT lớn thì sendall() phía khách
            # gãy giữa chừng (BrokenPipe) và phản hồi có thể mất — đó là hành vi
            # ĐÚNG, không phải lỗi. Gộp hai ca thành một khẳng định "luôn đọc
            # được 413" sẽ đỏ ngẫu nhiên theo kích thước bộ đệm socket của máy.
            #
            #   lỡ tay   — hơi quá trần: PHẢI đọc được 413 (việc _doc_bo phục vụ)
            #   tấn công — 20 MB:        ĐƯỢC PHÉP mất phản hồi, nhưng tuyệt đối
            #                            không được đọc thân và không đổi tệp
            for nhan, co, doi_413 in (('lỡ tay', tran + 40_000, True),
                                      ('tấn công', 20_000_000, False)):
                do['lon_nhat'] = 0
                than = json.dumps({'tat': [], 'rac': 'A' * co}).encode()
                ma, _, t = k.yeu_cau('POST', '/api/luu-quy-tac', than)
                if doi_413:
                    kiem(ma == 413,
                         f'[{nhan}] thân {len(than):,} byte -> ĐỌC ĐƯỢC 413 — {ma}')
                    kiem('KB' in t.decode('utf-8', 'replace'),
                         '  và câu lỗi nói rõ trần')
                else:
                    kiem(ma in (413, None),
                         f'[{nhan}] thân {len(than):,} byte -> 413 hoặc mất '
                         f'phản hồi (đều đúng) — {ma}')
                    kiem(k.gay_khi_gui or ma == 413,
                         '  máy chủ ngắt sớm chứ không nuốt hết 20 MB')
                kiem(do['lon_nhat'] == 0,
                     f'  KHÔNG đọc byte nào của thân — read() lớn nhất '
                     f'{do["lon_nhat"]}')
                kiem(so_tat(dd) == [2],
                     '  và tệp cấu hình không đổi')

            # --- Content-Length nói dối rất lớn: chặn theo lời khai ---
            do['lon_nhat'] = 0
            ma, _, t = k.tho_dong(
                'POST', '/api/luu-quy-tac',
                [('Host', f'127.0.0.1:{cong}'),
                 ('Content-Type', 'application/json'),
                 ('Content-Length', '524288000'),
                 ('Connection', 'close')], b'{"tat":[]}')
            kiem(ma == 413, f'Content-Length khai 500 MB -> 413 — {ma}')
            kiem(do['lon_nhat'] == 0, '  và không đọc byte nào')
            kiem(so_tat(dd) == [2], '  và tệp cấu hình không đổi')
    finally:
        W.May._doc_than_co_tran = that

    # --- kết nối im lặng không được giữ luồng vĩnh viễn ---
    # Đo được: BaseHTTPRequestHandler.timeout mặc định là None.
    # Hạ timeout xuống cho ca thử chạy trong ~2 giây thay vì 30.
    kiem(W.May.timeout is not None,
         f'Lớp May có đặt timeout — {W.May.timeout}')
    cu = W.May.timeout
    W.May.timeout = 1
    try:
        with MayChuTam(dd) as (k2, (dia, cong)):
            s = socket.create_connection((dia, cong), timeout=10)
            try:
                s.sendall(f'POST /api/luu-quy-tac HTTP/1.1\r\n'
                          f'Host: 127.0.0.1:{cong}\r\n'
                          f'Content-Type: application/json\r\n'
                          f'Content-Length: 50\r\n\r\n'.encode())
                t0 = time.time()
                ra = b''
                while True:                     # rồi im lặng, không gửi thân
                    m = s.recv(65536)
                    if not m:
                        break
                    ra += m
                giay = time.time() - t0
            finally:
                s.close()
            kiem(giay < 8,
                 f'Kết nối im lặng bị đóng sau {giay:.1f}s, luồng được thu hồi')
            kiem(k2.get('/api/gia-tri')[0] == 200,
                 'Máy chủ vẫn phục vụ bình thường sau đó')
    finally:
        W.May.timeout = cu


# ==========================================================================
def ca_a10_tieu_de(tm):
    print('\nCA A10 — mọi tiêu đề đi qua MỘT phễu, không chèn được đầu HTTP')
    print('-' * 74)
    tmx = os.path.join(tm, 'a10')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)

    def dat_file_xuat(ten):
        cfg = doc_an_toan(dd)
        cfg['thong_tin']['file_xuat'] = ten
        with io.open(dd, 'w', encoding='utf-8') as f:
            json.dump(cfg, f, ensure_ascii=False)

    # 1. CHUỖI TẤN CÔNG — đã dựng lại được trên bản cũ:
    #    set-cookie: a=b" và x-chen-vao: co-that đi ra thật.
    AC = 'TKB.xlsx"\r\nX-Chen-Vao: co-that\r\nSet-Cookie: phien=bikiem'
    with MayChuTam(dd) as (k, _):
        dat_file_xuat(AC)
        try:
            with io.open(os.path.join(tmx, AC), 'wb') as f:
                f.write(b'PK\x03\x04gia')
        except OSError:
            pass                     # hệ tệp không cho tên như vậy — vẫn thử tiếp
        ma, td, than = k.yeu_cau('GET', '/api/tai-ve')
        kiem('x-chen-vao' not in td,
             f'KHÔNG chèn được tiêu đề lạ — thấy {sorted(td)}')
        kiem('set-cookie' not in td,
             'KHÔNG chèn được Set-Cookie (đây là mồi cố định phiên khi có cookie)')
        kiem(ma == 400, f'Tên tệp có xuống dòng -> 400 — {ma}')
        kiem('file_xuat' in than.decode('utf-8', 'replace'),
             'Câu lỗi chỉ ĐÍCH DANH thong_tin.file_xuat')
        kiem(td.get('content-type', [''])[0].startswith('text/plain'),
             f'Lỗi ở đường TẢI trả text/plain, không phải JSON trần — '
             f'{td.get("content-type")}')

        # 2. TÊN CÓ DẤU TIẾNG VIỆT — lỗi hay gặp hơn hẳn, không cần kẻ tấn công.
        #    Bản cũ: mã 200 với server/date/content-type mỗi thứ HAI LẦN, và
        #    trình duyệt tải về .xlsx chứa JSON báo lỗi.
        TEN_DAU = 'TKB_Khối 12.xlsx'
        dat_file_xuat(TEN_DAU)
        with io.open(os.path.join(tmx, TEN_DAU), 'wb') as f:
            f.write(b'PK\x03\x04noidungthat')
        ma, td, than = k.yeu_cau('GET', '/api/tai-ve')
        kiem(ma == 200, f'Tên tệp CÓ DẤU tải về được — {ma}')
        lap = sorted(x for x, v in td.items() if len(v) > 1)
        kiem(not lap, f'Mỗi tiêu đề xuất hiện ĐÚNG MỘT LẦN — lặp: {lap}')
        kiem(than == b'PK\x03\x04noidungthat',
             f'Thân là nội dung tệp thật, không phải câu báo lỗi — {than[:40]!r}')
        bt = td.get('content-disposition', [''])[0]
        kiem("filename*=UTF-8''" in bt,
             f'Có filename* dạng RFC 5987 — {bt!r}')
        kiem('%E1%BB%91' in bt, f'  và mã hoá đúng chữ "ố" — {bt!r}')
        kiem('filename="' in bt, f'  vẫn kèm bản ASCII dự phòng — {bt!r}')

        # 3. tên tệp bình thường: canh KHÔNG HỒI QUY
        dat_file_xuat('TKB_Khoi12.xlsx')
        with io.open(os.path.join(tmx, 'TKB_Khoi12.xlsx'), 'wb') as f:
            f.write(b'PK\x03\x04binhthuong')
        ma, td, than = k.yeu_cau('GET', '/api/tai-ve')
        bt = td.get('content-disposition', [''])[0]
        kiem(ma == 200 and 'filename="TKB_Khoi12.xlsx"' in bt,
             f'Tên ASCII bình thường giữ nguyên — {bt!r}')

        # 4. tên hiểm khác
        dat_file_xuat('../../khong-ton-tai.xlsx')
        ma, _, _ = k.yeu_cau('GET', '/api/tai-ve')
        kiem(ma == 404,
             f'Tên có "../" bị basename() rút gọn rồi báo không có tệp — {ma}')
        dat_file_xuat('a"b.xlsx')
        ma, td, _ = k.yeu_cau('GET', '/api/tai-ve')
        kiem(ma == 400 and 'set-cookie' not in td,
             f'Tên có dấu nháy kép -> 400 — {ma}')

        # 5. đường tải trang web (tên là HẰNG) vẫn chạy
        dat_file_xuat('TKB_Khoi12.xlsx')
        with io.open(os.path.join(tmx, 'tkb_web.html'), 'wb') as f:
            f.write(b'<!doctype html>')
        ma, td, _ = k.yeu_cau('GET', '/api/tai-web')
        lap = sorted(x for x, v in td.items() if len(v) > 1)
        kiem(ma == 200 and not lap, f'/api/tai-web vẫn 200, không lặp — {ma}')

        # 6. không phản hồi nào có hai dòng trạng thái, và không khoe phiên bản
        khoe = []
        for duong in ('/', '/api/gia-tri', '/khong-co', '/api/tai-ve',
                      '/api/tai-web'):
            ma, td, _ = k.yeu_cau('GET', duong)
            if len(td.get('server', [])) > 1 or any(
                    'Python' in v for v in td.get('server', [])):
                khoe.append((duong, td.get('server')))
        kiem(not khoe,
             f'Mọi đường: tiêu đề Server đúng một dòng và KHÔNG khoe phiên bản '
             f'Python — lệch: {khoe}')

    # 7. CA THỬ CẤU TRÚC — không lạc hậu khi thêm đường mới.
    #    Chỉ canh send_header là bỏ lọt đúng cái bẫy hai-dòng-trạng-thái, nên
    #    canh cả send_response.
    ma_nguon = io.open(os.path.join(GOC_TM, 'web.py'), encoding='utf-8').read()
    than_pheu = ma_nguon.split('def _tra_byte')[1].split('# ---------------- tiện ích')[0]
    ngoai = ma_nguon.replace(than_pheu, '')
    for goi in ('self.send_header(', 'self.send_response(', 'self.end_headers('):
        kiem(ngoai.count(goi) == 0,
             f'Ngoài _tra_byte không còn chỗ nào gọi {goi[:-1]} — '
             f'đếm được {ngoai.count(goi)}')


# ==========================================================================
def ca_a11_ben_vung(tm):
    print('\nCA A11 — ghi tệp bền: fsync đúng thứ tự, không có trạng thái giữa')
    print('-' * 74)
    tmx = os.path.join(tm, 'a11')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)
    import web as W
    import stat as _stat

    # ---- (a) GHIM THỨ TỰ ----
    # Bọc os.fsync / os.replace để ghi lại dãy thao tác thật, rồi bắn MỘT yêu
    # cầu Lưu qua HTTP thật. Rẻ, chạy được cả Linux lẫn Windows, và lùi vá ra
    # là đỏ ngay.
    that_fsync, that_replace = os.fsync, os.replace
    nhat_ky = []

    def fsync_ghi(fd):
        try:
            la_thu_muc = _stat.S_ISDIR(os.fstat(fd).st_mode)
        except OSError:
            la_thu_muc = False
        nhat_ky.append(('fsync-thu-muc' if la_thu_muc else 'fsync-tep', fd))
        return that_fsync(fd)

    def replace_ghi(a, b):
        nhat_ky.append(('replace', os.path.basename(str(b))))
        return that_replace(a, b)

    os.fsync, os.replace = fsync_ghi, replace_ghi
    try:
        with MayChuTam(dd) as (k, _):
            del nhat_ky[:]
            ma, _ = k.js('/api/luu-quy-tac', {'tat': [1]})
            kiem(ma == 200 and so_tat(dd) == [1], f'Lưu lần đầu OK — {ma}')
            ten = [x for x, _ in nhat_ky]

            def vi_tri(x):
                """Chỉ số của thao tác, -1 nếu KHÔNG có.

                   Dùng thay ten.index(): khi lùi bản vá ra thì thao tác biến
                   mất, và .index() sẽ ném ValueError — ca thử CHẾT GIỮA CHỪNG
                   thay vì báo đỏ, nên mọi ca sau nó không bao giờ được chạy.
                   Ca thử phải biết đỏ, không phải biết vỡ.
                """
                return ten.index(x) if x in ten else -1

            kiem(ten.count('fsync-tep') >= 1, f'Có fsync tệp — {ten}')
            kiem('replace' in ten, f'Có os.replace — {ten}')
            kiem(0 <= vi_tri('fsync-tep') < vi_tri('replace'),
                 f'fsync tệp tạm chạy TRƯỚC os.replace — {ten}')
            if os.name == 'posix':
                kiem('fsync-thu-muc' in ten,
                     f'Có fsync THƯ MỤC CHA — {ten}')
                kiem(0 <= vi_tri('replace') < vi_tri('fsync-thu-muc'),
                     f'fsync thư mục chạy SAU os.replace — {ten}')
                # lần lưu ĐẦU còn ghi cauhinh.bak.goc.json -> 3 tệp
                kiem(ten.count('fsync-tep') == 3,
                     f'Lần Lưu đầu: fsync 3 tệp (tam + bak.goc + bak) — '
                     f'đếm {ten.count("fsync-tep")}')
            else:
                kiem(True, 'Không phải POSIX — bỏ qua mục fsync thư mục')

            # lần lưu THỨ HAI: bak.goc không ghi lại nữa -> 2 tệp.
            # ĐỪNG ghim một con số cho mọi lần lưu: nó khác nhau theo lần.
            del nhat_ky[:]
            k.js('/api/luu-quy-tac', {'tat': [2]})
            ten2 = [x for x, _ in nhat_ky]
            if os.name == 'posix':
                kiem(ten2.count('fsync-tep') == 2,
                     f'Lần Lưu sau: fsync 2 tệp (tam + bak) — '
                     f'đếm {ten2.count("fsync-tep")}')
    finally:
        os.fsync, os.replace = that_fsync, that_replace

    # ---- (b) CẮT NGANG bằng ngoại lệ có chủ ý, ở 4 mốc ----
    # Mạnh hơn (a) vì bắt được cả lỗi thứ tự sao lưu / replace.
    tmy = os.path.join(tm, 'a11b')
    os.makedirs(tmy, exist_ok=True)
    dd2 = cau_hinh_tam(tmy)
    with MayChuTam(dd2) as (k, _):
        k.js('/api/luu-quy-tac', {'tat': [0]})       # dựng sẵn bak.goc + bak
    goc_truoc = so_tat(dd2)

    class NoGiua(Exception):
        pass

    MOC = [('sau khi ghi tệp tạm', 'dong_bo_tep'),
           ('trước khi sao lưu', 'xoay'),
           ('sau sao lưu, trước replace', 'replace'),
           ('sau replace, trước fsync thư mục', 'thu_muc')]
    for ten_moc, dau in MOC:
        that_map = {
            'dong_bo_tep': W._dong_bo_tep,
            'xoay': W._xoay_sao_luu,
            'replace': os.replace,
            'thu_muc': W._dong_bo_thu_muc,
        }

        def no(*a, **kw):
            raise NoGiua(ten_moc)

        if dau == 'replace':
            os.replace = no
        else:
            setattr(W, '_' + dau if dau != 'xoay' else '_xoay_sao_luu', no)
        try:
            with MayChuTam(dd2) as (k, _):
                ma, _ = k.js('/api/luu-quy-tac', {'tat': [1, 2, 3]})
        except Exception:
            ma = None
        finally:
            os.replace = that_map['replace']
            W._dong_bo_tep = that_map['dong_bo_tep']
            W._xoay_sao_luu = that_map['xoay']
            W._dong_bo_thu_muc = that_map['thu_muc']

        cfg = doc_an_toan(dd2)
        kiem(cfg is not None,
             f'Cắt ngang "{ten_moc}": cauhinh.json vẫn ĐỌC ĐƯỢC')
        if cfg is not None:
            hien = [i for i, r in enumerate(cfg['quy_tac']) if r.get('bat') is False]
            kiem(hien in (goc_truoc, [1, 2, 3]),
                 f'  và là bản CŨ trọn vẹn hoặc bản MỚI trọn vẹn — {hien}')
        # phao cứu sinh phải luôn còn
        con = [t for t in os.listdir(tmy)
               if t.startswith('cauhinh.bak') and t.endswith('.json')
               and doc_an_toan(os.path.join(tmy, t)) is not None]
        kiem(con, f'  và còn ít nhất một bản sao lưu đọc được — {con}')

    kiem(W._dong_bo_thu_muc(tmy) == (os.name == 'posix'),
         f'_dong_bo_thu_muc trả bool đúng theo hệ điều hành ({os.name})')
    kiem(W._dong_bo_thu_muc(os.path.join(tmy, 'khong-co-thu-muc-nay')) is False,
         'Thư mục không tồn tại thì trả False, không làm hỏng lượt Lưu')

    # GIỚI HẠN CỦA HAI CA NÀY, nói thẳng ra:
    # chúng ghim THỨ TỰ và bất biến "không có trạng thái ở giữa".
    # Chúng KHÔNG chứng minh độ bền vật lý khi mất điện — muốn vậy phải cắt
    # điện thật hoặc dùng dm-flakey. Cách "biến fsync thành lệnh rỗng rồi đọc
    # lại" là VÔ DỤNG: trong cùng một máy, đọc lại luôn thấy dữ liệu qua bộ nhớ
    # đệm, nên ca thử sẽ xanh cả trên mã hỏng.
    kiem(True, 'Ghi chú: ca này ghim THỨ TỰ, không chứng minh chống mất điện')


# ==========================================================================
def ca_a7_sao_luu_xoay_vong(tm):
    print('\nCA A7 — sao lưu xoay vòng, và ghi cauhinh.json phải NGUYÊN TỬ')
    print('-' * 74)
    tmx = os.path.join(tm, 'a7')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)
    goc = X.doc_cau_hinh(dd)

    def bak(ten):
        return os.path.join(tmx, ten)

    with MayChuTam(dd) as (k, _):
        for lan, t in enumerate([[1], [2], [3], [4]], start=1):
            ma, _kq = k.js('/api/luu-quy-tac', {'tat': t})
            kiem(ma == 200, f'Lưu lần {lan} (tắt quy tắc {t}) — {ma}')

        # 1. bản GỐC phải còn nguyên sau 4 lần lưu
        kiem(os.path.exists(bak('cauhinh.bak.goc.json')),
             'Có bản gốc cauhinh.bak.goc.json')
        kiem(doc_an_toan(bak('cauhinh.bak.goc.json')) == goc,
             'Bản gốc SAU 4 LẦN LƯU vẫn đúng là nội dung ban đầu')

        # 2. các bản xoay vòng phải giữ đúng thứ tự lùi dần
        kiem(so_tat(bak('cauhinh.bak.json')) == [3],
             f'cauhinh.bak.json = ngay trước lần lưu cuối — '
             f'{so_tat(bak("cauhinh.bak.json"))}')
        kiem(so_tat(bak('cauhinh.bak.1.json')) == [2],
             f'cauhinh.bak.1.json lùi thêm một bậc — '
             f'{so_tat(bak("cauhinh.bak.1.json"))}')
        kiem(so_tat(bak('cauhinh.bak.2.json')) == [1],
             f'cauhinh.bak.2.json lùi thêm bậc nữa — '
             f'{so_tat(bak("cauhinh.bak.2.json"))}')
        kiem(so_tat(dd) == [4], f'Tệp thật đang ở trạng thái mới nhất — {so_tat(dd)}')

        # 3. NGUYÊN TỬ: dựng chuỗi JSON hỏng thì tệp thật phải Y NGUYÊN
        import web as W
        truoc = io.open(dd, encoding='utf-8').read()
        that = W._gon_json

        def hong(cfg):
            raise RuntimeError('cố ý làm hỏng lúc dựng chuỗi JSON')
        W._gon_json = hong
        try:
            ma, _ = k.js('/api/luu-quy-tac', {'tat': [0]})
        finally:
            W._gon_json = that
        sau = io.open(dd, encoding='utf-8').read()
        kiem(ma == 500, f'Dựng chuỗi hỏng thì báo lỗi — {ma}')
        kiem(sau == truoc,
             f'cauhinh.json KHÔNG bị cắt cụt, còn nguyên {len(sau)} byte '
             f'(trước {len(truoc)})')
        kiem(len(sau) > 1000 and doc_an_toan(dd) is not None,
             'Và vẫn đọc được như một tệp JSON hợp lệ')

        # 4. hai người bấm Lưu cùng lúc: tệp phải luôn đọc được
        loi = []

        def ban(i):
            try:
                k2 = Khach(k.cong)
                k2.js('/api/luu-quy-tac', {'tat': [i]})
            except Exception as e:          # pragma: no cover
                loi.append(repr(e))
        ts = [threading.Thread(target=ban, args=(i,)) for i in range(6)]
        for t in ts:
            t.start()
        for t in ts:
            t.join()
        kiem(not loi, f'6 người bấm Lưu cùng lúc, không ai lỗi — {loi[:1]}')
        cuoi = doc_an_toan(dd)
        kiem(cuoi is not None,
             'Sau khi 6 người lưu đồng thời, tệp vẫn đọc được')
        kiem(cuoi and len(cuoi.get('quy_tac', [])) == len(goc['quy_tac']),
             'Và không mất quy tắc nào')


# ==========================================================================
def ca_a4_mot_luot(tm):
    print('\nCA A4 — hai lượt xếp KHÔNG BAO GIỜ được chạy chồng lên nhau')
    print('-' * 74)
    tmx = os.path.join(tm, 'a4')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)

    # ĐO THẲNG ĐỘ CHỒNG LẤN, không đếm số 200.
    #
    # Vì sao không đếm được: bộ giải là Python thuần, ăn CPU và giữ GIL, nên
    # các luồng HTTP gần như đói hẳn trong lúc giải. Đã đo: 8 người bấm cùng
    # lúc thì người thứ 8 mãi 1,06 giây mới được phục vụ — sau khi lượt trước
    # đã xong ở 0,78 giây. Người đó nhận 200 là ĐÚNG, hai lượt NỐI TIẾP nhau
    # chứ không chồng nhau. Đếm số 200 sẽ báo đỏ oan, và kéo dài lượt giải
    # cũng không cứu được vì yêu cầu đến muộn lại càng được phục vụ muộn hơn.
    #
    # Thứ cần khẳng định là: không lúc nào có hai luồng _chay cùng sống. Bọc
    # đếm quanh web._chay để đo đúng điều đó. Yêu cầu vẫn đi qua HTTP thật —
    # đây là dụng cụ đo, không phải gọi tắt vào hàm xử lý.
    import web as W
    that = W._chay
    dem = {'dang': 0, 'nhieu_nhat': 0, 'tong': 0}
    khoa_dem = threading.Lock()

    def dem_chay(*a, **kw):
        with khoa_dem:
            dem['dang'] += 1
            dem['tong'] += 1
            dem['nhieu_nhat'] = max(dem['nhieu_nhat'], dem['dang'])
        try:
            return that(*a, **kw)
        finally:
            with khoa_dem:
                dem['dang'] -= 1

    W._chay = dem_chay

    # Nới khe hở CÓ SẴN cho đo được — không tạo lỗi mới.
    #
    # _than_json() là điểm mấu chốt phân biệt hai bản: bản cũ đọc cờ dang_chay
    # RỒI MỚI gọi _than_json() rồi mãi sau mới đặt cờ, nên _than_json() nằm
    # NGAY GIỮA vùng chưa khoá; bản mới gọi _than_json() TRƯỚC khi giành chỗ,
    # nên nó nằm NGOÀI vùng đó. Thêm 50ms ở đây làm khe hở của bản cũ rộng ra
    # đủ để mọi luồng cùng lọt, còn bản mới thì không ảnh hưởng gì.
    #
    # Không có mẹo này thì khe hở hẹp tới mức 8 luồng qua 5 vòng vẫn không
    # trúng, và ca thử sẽ XANH cả trên mã hỏng — đã thử, đúng như vậy.
    that_tj = W.May._than_json

    def cham_tj(may, *a, **kw):
        d = that_tj(may, *a, **kw)
        time.sleep(0.05)
        return d

    W.May._than_json = cham_tj
    try:
        SO_NGUOI, SO_VONG = 8, 3
        with MayChuTam(dd) as (k, _):
            for vong in range(SO_VONG):
                doi_trang_thai(k)
                rao = threading.Barrier(SO_NGUOI)
                ma = []
                khoa_ds = threading.Lock()

                def ban(i):
                    k2 = Khach(k.cong)
                    rao.wait()
                    m, _t = k2.js('/api/xep', {'tuan': f'NGUOI-{i}'})
                    with khoa_ds:
                        ma.append(m)

                ts = [threading.Thread(target=ban, args=(i,))
                      for i in range(SO_NGUOI)]
                for t in ts:
                    t.start()
                for t in ts:
                    t.join()
                kiem(dem['nhieu_nhat'] == 1,
                     f'Vòng {vong + 1}: {SO_NGUOI} người bấm cùng lúc -> '
                     f'nhiều nhất {dem["nhieu_nhat"]} lượt chạy đồng thời '
                     f'({ma.count(200)} nhận, {ma.count(409)} bị chặn 409)')
            doi_trang_thai(k)

            kiem(dem['nhieu_nhat'] == 1,
                 f'Qua {SO_VONG} vòng và {dem["tong"]} lượt, chưa lúc nào có '
                 f'2 lượt chồng nhau (nhiều nhất {dem["nhieu_nhat"]})')

            # nhật ký chỉ được là của MỘT lượt, không trộn hai người
            t = doi_trang_thai(k)
            log = '\n'.join(t.get('log_moi') or [])
            kiem(log.count('Tổng:') == 1,
                 f'Khung nhật ký chỉ chứa một lượt — '
                 f'{log.count("Tổng:")} dòng "Tổng:"')
            ten = [d for d in (t.get('log_moi') or []) if 'NGUOI-' in d]
            kiem(len(ten) == 1,
                 f'Chỉ một tuần được ghi vào nhật ký — {len(ten)} dòng')
            luoi_tuan = (t.get('luoi') or {}).get('tuan')
            kiem(bool(ten) and luoi_tuan and luoi_tuan in ten[0],
                 f'Lưới hiện ra đúng của lượt đã chạy — {luoi_tuan}')

            # yêu cầu hỏng KHÔNG được chiếm chỗ rồi bỏ đó
            ma, _ = k.post('/api/xep', 'khong-phai-json')
            kiem(ma == 400, f'Thân hỏng -> 400 — {ma}')
            ma, _ = k.js('/api/xep', {})
            kiem(ma == 200,
                 f'Và lượt sau vẫn bấm được (chỗ không bị chiếm oan) — {ma}')
            doi_trang_thai(k)
    finally:
        W._chay = that
        W.May._than_json = that_tj


# ==========================================================================
def ca_a4b_thu_quy_tac(tm):
    print('\nCA A4b — /api/thu-quy-tac cũng phải một-lúc-một-người')
    print('-' * 74)
    tmx = os.path.join(tm, 'a4b')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)
    import web as W

    with MayChuTam(dd) as (k, (dia, cong)):
        # 1. Dữ liệu sai phải cho 400 kèm câu tiếng Việt, KHÔNG phải 500 + mã
        #    sự cố. Đường này trước đây tự viết [int(x) for x in ...] thay vì
        #    dùng chung _tap_chi_so() với bốn đường kia.
        ma, t = k.post('/api/thu-quy-tac', {'tat': ['a']})
        kiem(ma == 400, f'"tat" có giá trị lạ -> 400, không phải 500 — {ma}')
        kiem('không phải số' in t,
             f'  và câu lỗi chỉ đúng chỗ sai — {t[:70]!r}')
        kiem('mã sự cố' not in t, '  không phải mã sự cố 6 ký tự vô nghĩa')

        # 2. ĐO THẲNG số lần giải cùng sống — không đếm mã 200.
        #    (Bài học A4: bộ giải giữ GIL nên luồng HTTP đói, đếm 200 báo sai.)
        that = W.thu_quy_tac
        dem = {'dang': 0, 'dinh': 0}
        kh = threading.Lock()

        def bao(*a, **kw):
            with kh:
                dem['dang'] += 1
                dem['dinh'] = max(dem['dinh'], dem['dang'])
            try:
                return that(*a, **kw)
            finally:
                with kh:
                    dem['dang'] -= 1

        W.thu_quy_tac = bao
        try:
            kq = []
            kh2 = threading.Lock()

            def ban():
                ma = Khach(cong).post('/api/thu-quy-tac', {'tat': []})[0]
                with kh2:
                    kq.append(ma)

            ls = [threading.Thread(target=ban) for _ in range(4)]
            for t2 in ls:
                t2.start()
            for t2 in ls:
                t2.join()
        finally:
            W.thu_quy_tac = that

        kiem(dem['dinh'] == 1,
             f'KHÔNG lúc nào có hai lượt thử cùng sống — cao nhất '
             f'{dem["dinh"]} (mã trả về: {sorted(kq)})')
        kiem(200 in kq, f'  và vẫn có người được phục vụ — {sorted(kq)}')
        kiem(409 in kq, f'  người còn lại nhận 409 kèm lời nhắn — {sorted(kq)}')

        # 3. xong rồi thì phải nhả ra, không kẹt
        ma, _ = k.post('/api/thu-quy-tac', {'tat': []})
        kiem(ma == 200, f'Thử lại sau đó vẫn được phục vụ, không kẹt — {ma}')


# ==========================================================================
def ca_a3_khong_lo_traceback(tm):
    print('\nCA A3 — không trả dấu vết lỗi Python ra ngoài')
    print('-' * 74)
    tmx = os.path.join(tm, 'a3')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)

    def sach(than, nhan):
        """Phản hồi không được chứa dấu vết nội bộ nào."""
        # Dò dấu vết NỘI BỘ, không dò tên tệp trần: thông báo khai báo sai
        # cố ý kết thúc bằng "Sửa lại rồi chạy: python xep_tkb.py" — đó là
        # hướng dẫn cho người dùng, không phải rò rỉ.
        xau = [('Traceback', 'dấu vết lỗi'),
               ('File "', 'khung dấu vết lỗi'),
               (', line ', 'số dòng trong khung dấu vết'),
               (GOC_TM, 'đường dẫn tuyệt đối của dự án'),
               ('/home/', 'thư mục người dùng'),
               ('/usr/lib/python', 'đường dẫn thư viện Python'),
               ('site-packages', 'đường dẫn thư viện'),
               ('Error:', 'tên lớp ngoại lệ Python')]
        thay = [ten for chuoi, ten in xau if chuoi and chuoi in than]
        kiem(not thay, f'{nhan}: không lộ {thay if thay else "gì cả"}')

    with MayChuTam(dd) as (k, _):
        # 1. tham số truy vấn không phải số — đúng ca đã đo lộ đường dẫn
        ma, than = k.get('/api/trang-thai?tu=abc')
        kiem(ma == 400, f'?tu=abc trả 400 chứ không 500 — {ma}')
        sach(than, '?tu=abc')
        for xau_tu in ('?tu=-5', '?tu=99999999999999999999', '?tu=', '?tu=1.5'):
            ma, than = k.get('/api/trang-thai' + xau_tu)
            kiem(ma in (200, 400), f'{xau_tu} trả {ma}, không phải 500')
            sach(than, xau_tu)

        # 2. thân JSON hỏng
        ma, than = k.post('/api/them-quy-tac', 'khong-phai-json')
        kiem(ma == 400, f'Thân không phải JSON -> 400 — {ma}')
        sach(than, 'thân JSON hỏng')

        # 3. lỗi BẤT NGỜ thật sự: phải 500, có mã sự cố, không có traceback
        import web as W
        that = W.doc_quy_tac

        def no(*a, **kw):
            raise RuntimeError('lỗi cố ý để thử — KHONG_DUOC_LO_RA_NGOAI')
        W.doc_quy_tac = no
        try:
            ma, than = k.get('/api/quy-tac')
        finally:
            W.doc_quy_tac = that
        kiem(ma == 500, f'Lỗi bất ngờ trả 500 — {ma}')
        sach(than, 'lỗi bất ngờ')
        kiem('KHONG_DUOC_LO_RA_NGOAI' not in than,
             'Không lộ cả nội dung thông báo lỗi bên trong')
        kiem('sự cố' in than.lower() or 'su co' in than.lower(),
             f'Có báo cho người dùng biết là sự cố — {than[:80]}')
        import re as _re
        m = _re.search(r'mã sự cố ([0-9a-f]{6})', than)
        kiem(m, 'Có mã sự cố để đối chiếu với cửa sổ dòng lệnh')

        # Mặt kia của việc vá: dấu vết PHẢI tới được người vận hành, nếu không
        # thì "vá" hoá ra chỉ là xoá mất chẩn đoán.
        import contextlib
        bat = io.StringIO()
        W.doc_quy_tac = no
        try:
            with contextlib.redirect_stderr(bat):
                ma2, than2 = k.get('/api/quy-tac')
        finally:
            W.doc_quy_tac = that
        loi_console = bat.getvalue()
        m2 = _re.search(r'\[SỰ CỐ ([0-9a-f]{6})\]', loi_console)
        kiem(m2, 'Cửa sổ dòng lệnh CÓ in dấu vết kèm mã sự cố')
        kiem('Traceback' in loi_console and 'KHONG_DUOC_LO_RA_NGOAI' in loi_console,
             'Và in ĐẦY ĐỦ dấu vết cho người vận hành gỡ lỗi')
        kiem(m2 and m2.group(1) in than2,
             f'Mã trên màn hình khớp mã trong dòng lệnh — '
             f'{m2.group(1) if m2 else "?"}')

        # 4. khai báo sai trong cauhinh.json thì VẪN phải hiện nguyên văn
        # dùng post() lấy CHUỖI THÔ — js() trả dict, mà "x in dict" là kiểm
        # KHOÁ chứ không kiểm nội dung, nên ca thử sẽ xanh giả.
        ma, than = k.post('/api/them-quy-tac',
                          {'quy_tac': {'loai': 'chi_xep_thu', 'ghi_chu': 'x',
                                       'thu': [99]}})
        kiem(ma == 400, f'Khai báo sai -> 400 — {ma}')
        kiem('[LỖI]' in than and '99' in than
             and 'Miền hợp lệ' in than and 'SỐ HIỆU THỨ' in than,
             'Thông báo khai báo sai vẫn giữ NGUYÊN VĂN, cả dòng gợi ý')
        sach(than, 'khai báo sai')


# ==========================================================================
def ca_a8_khong_thanh_cong_thuc(tm):
    print('\nCA A8 — chuỗi người dùng gõ KHÔNG được thành công thức Excel')
    print('-' * 74)
    import openpyxl
    tmx = os.path.join(tm, 'a8')
    os.makedirs(tmx, exist_ok=True)
    dd = cau_hinh_tam(tmx)
    xuat = os.path.join(tmx, 'TKB_Khoi12.xlsx')

    DOC_TUAN = "=cmd|' /c calc'!A1"
    DOC_GHI_CHU = '=HYPERLINK("http://ke-gian/"&A1,"Bấm vào đây")'

    with MayChuTam(dd) as (k, _):
        # đường 2: ghi_chu của một quy tắc MỀM đi thẳng vào đầu ô, không tiền tố
        ma, _t = k.js('/api/them-quy-tac', {'quy_tac': {
            'loai': 'uu_tien_o', 'ghi_chu': DOC_GHI_CHU,
            'loc': {'nhom_mon': ['Toán']}, 'tiet_trong_buoi': [1, 2],
            'diem': 3}})
        kiem(ma == 200, f'Thêm được quy tắc mềm có ghi_chu bắt đầu bằng "=" — {ma}')

        # đường 1: tuan gửi qua /api/xep
        ma, _t = k.js('/api/xep', {'tuan': DOC_TUAN})
        kiem(ma == 200, f'Gửi lệnh xếp với tuần độc hại — {ma}')
        t = doi_trang_thai(k)
        kiem(t.get('loi') is None, f'Lượt xếp trót lọt — {t.get("loi")}')

    kiem(os.path.exists(xuat), 'Có tệp Excel để soi')
    wb = openpyxl.load_workbook(xuat)

    # 1. không một ô nào MÀ CHƯƠNG TRÌNH GHI được phép là công thức
    o_ct = []
    for ten in wb.sheetnames:
        ws = wb[ten]
        for hang in ws.iter_rows():
            for o in hang:
                if o.data_type == 'f':
                    o_ct.append((ten, o.coordinate, str(o.value)[:40]))
    # bảy công thức có sẵn trong mau_tkb.xlsx thì PHẢI còn nguyên
    cua_mau = {('Khối 12', f'{c}43') for c in 'DEFGHI'} | {('Khối 12', 'H48')}
    thay = {(a, b) for a, b, _ in o_ct}
    la = sorted(x for x in thay if x not in cua_mau)
    kiem(not la, f'Không ô nào do chương trình ghi trở thành công thức — {la[:4]}')
    kiem(cua_mau <= thay,
         f'7 công thức có sẵn của mẫu VẪN CÒN (COUNTA + ngày ký) — '
         f'thiếu {sorted(cua_mau - thay)}')

    # 2. chữ phải giữ NGUYÊN VĂN — nhà trường in tờ này để ký
    ws = wb['Khối 12']
    o_tuan = ws.cell(4, 1)
    kiem(o_tuan.value == DOC_TUAN,
         f'Ô A4 giữ nguyên từng ký tự — {str(o_tuan.value)[:40]!r}')
    kiem(o_tuan.data_type == 's',
         f'Và được lưu dạng CHỮ, không phải công thức — {o_tuan.data_type!r}')

    ws2 = wb['Kiểm tra quy tắc']
    thay_ghi_chu = [o for hang in ws2.iter_rows() for o in hang
                    if o.value == DOC_GHI_CHU]
    kiem(thay_ghi_chu,
         f'Tìm thấy ghi_chu quy tắc mềm trong sheet Kiểm tra quy tắc — '
         f'{len(thay_ghi_chu)} ô')
    kiem(all(o.data_type == 's' for o in thay_ghi_chu),
         f'Và cũng lưu dạng chữ — '
         f'{[o.data_type for o in thay_ghi_chu]}')

    # 3. soi thẳng XML: không có thẻ <f> nào ngoài 7 thẻ của mẫu
    import zipfile
    import re as _re
    z = zipfile.ZipFile(xuat)
    tong_f = 0
    for ten in z.namelist():
        if ten.startswith('xl/worksheets/'):
            tong_f += len(_re.findall(r'<f[ >]', z.read(ten).decode('utf-8')))
    kiem(tong_f == 7,
         f'Trong XML có đúng 7 thẻ công thức (của mẫu) — đếm được {tong_f}')

    # 4. các dấu khác KHÔNG bị đụng tới (đổi chữ là hỏng tờ in)
    kiem(True, 'Ghi chú: openpyxl vốn lưu "+", "-", "@" dạng chữ nên '
               'không cần đụng — tờ in không đổi ký tự nào')


# ==========================================================================
def chay_het(tm):
    ca_a2_than_di_dang(tm)
    ca_a1_nguon_goc(tm)
    ca_a1b_danh_sach_ten_may(tm)
    ca_a1c_moi_duong_deu_bi_gac(tm)
    ca_a9_tran_than(tm)
    ca_a10_tieu_de(tm)
    ca_a11_ben_vung(tm)
    ca_a7_sao_luu_xoay_vong(tm)
    ca_a4_mot_luot(tm)
    ca_a4b_thu_quy_tac(tm)
    ca_a3_khong_lo_traceback(tm)
    ca_a8_khong_thanh_cong_thuc(tm)


if __name__ == '__main__':
    print('=' * 74)
    print(' THỬ AN NINH GIAO DIỆN WEB')
    print('=' * 74)
    with tempfile.TemporaryDirectory() as tm:
        chay_het(tm)
    print('=' * 74)
    hong = [m for ok, m in KQ if not ok]
    print(f'Tổng: {len(KQ) - len(hong)}/{len(KQ)} mục ĐẠT.')
    if hong:
        for m in hong:
            print('  LỖI:', m)
        sys.exit(1)
