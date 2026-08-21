#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
==========================================================================
 CA THỬ: SÁU Ô "TỔNG SỐ TIẾT/TUẦN" PHẢI ĐẾM TỪ DÒNG ĐẦU VÙNG DỮ LIỆU
--------------------------------------------------------------------------
 Chạy:  python3 tests/thu_o_tong.py

 VÌ SAO CÓ TỆP NÀY. Ngày 21/8/2026 nhà trường đưa ra một tệp TKB tuần 1 làm
 bằng tay, dòng 43 khai:

     D43=COUNTA(D10:D42)   E43=COUNTA(E11:E42)   F43=COUNTA(F10:F42)
     G43=COUNTA(G8:G42)    H43=COUNTA(H8:H42)    I43=COUNTA(I8:I42)

 Dữ liệu bắt đầu từ dòng 8, nên ba cột đầu bỏ sót 2-3 dòng đầu bảng. Cả sáu
 lớp đều có 23 tiết mà bảng in ra 21/20/21/23/23/23. Không một dấu hiệu nào
 báo: ô vẫn ra số, số vẫn trông hợp lý, chỉ là SAI.

 ĐÃ ĐO 21/8/2026 — mau_tkb.xlsx CỦA DỰ ÁN KHÔNG DÍNH LỖI ĐÓ. Cả sáu ô đều
 là COUNTA(<cột>8:<cột>42), và tệp xuất ra đếm đúng. Tệp hỏng là tệp làm tay
 của trường, phần mềm không hề đụng tới. Vì vậy tệp ca thử này KHÔNG đi kèm
 một bản vá nào — nó là HÀNG RÀO giữ nguyên trạng đang đúng.

 VÌ SAO CẦN HÀNG RÀO KHI ĐANG ĐÚNG. tests/thu_an_ninh.py CA A8 đã canh dòng
 43, nhưng nó chỉ hỏi "bảy công thức CÒN hay MẤT" — nó KHÔNG đọc vùng bên
 trong. Đổi D8:D42 thành D10:D42 thì A8 vẫn xanh nguyên. Đã đo: xem CA 3 và CA 4.

 KHÔNG GHIM SỐ TRẦN. Dòng đầu vùng dữ liệu lấy từ thong_tin.vi_tri_bang, số
 dòng phải phủ suy ra từ khung_gio. Ghim "8" và "42" vào đây là ghim hằng số
 lấy từ cauhinh.json — đúng điều tests/thu_ghim_du_lieu.py cấm, và cũng làm
 ca thử vô nghĩa nếu mai này bảng dời chỗ.
==========================================================================
"""
import io
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEET_XML = 'xl/worksheets/sheet1.xml'

KQ = []


def kiem(dieu_kien, mo_ta):
    KQ.append((bool(dieu_kien), mo_ta))
    print(f'  [{"ĐẠT" if dieu_kien else "LỖI"}] {mo_ta}')


def dat_dong_ho_canh(giay=180):
    """Bộ thử không bao giờ được phép TREO — hết giờ thì ĐỎ. Xem CLAUDE.md."""
    def bao():
        time.sleep(giay)
        print(f'\n[HẾT GIỜ] Quá {giay} giây — coi như ĐỎ, thoát.', flush=True)
        os._exit(1)
    t = threading.Thread(target=bao, daemon=True)
    t.start()


# ==========================================================================
# BỘ DÒ — một chỗ duy nhất, CA 0 chứng minh nó biết bắt và biết bỏ qua
# ==========================================================================
# Bắt cả COUNTA(D8:D42) lẫn COUNTA($D$8:$D$42); không bắt COUNT/COUNTIF.
MAU_COUNTA = re.compile(
    r'^COUNTA\(\s*\$?([A-Z]{1,3})\$?(\d{1,7})\s*:\s*\$?([A-Z]{1,3})\$?(\d{1,7})\s*\)$',
    re.IGNORECASE)

MAU_O = re.compile(r'<c r="([A-Z]{1,3}\d{1,7})"[^>]*>\s*<f>([^<]*)</f>')


def doc_cong_thuc(duong_dan):
    """Trả dict {'D43': 'COUNTA(D8:D42)', ...}. Toàn phần: tệp hỏng -> {}.

    Đọc thẳng XML chứ không qua openpyxl: đây là tệp mẫu do nhà trường giữ,
    và phần cần canh chính là VĂN BẢN công thức, không phải giá trị đã tính.
    """
    try:
        with zipfile.ZipFile(duong_dan) as z:
            x = z.read(SHEET_XML).decode('utf-8', 'replace')
    except Exception:
        return {}
    return {o: ct for o, ct in MAU_O.findall(x)}


def kiem_vung(cong_thuc, cac_cot, hang_tong, hang_dau, hang_cuoi_ghi):
    """Trả danh sách lời phàn nàn. Rỗng = sạch.

    hang_dau      = dòng ĐẦU vùng dữ liệu (suy từ vi_tri_bang)
    hang_cuoi_ghi = dòng CUỐI mà chương trình còn ghi vào (suy từ khung_gio)
    """
    loi = []
    ket = []
    for cot in cac_cot:
        o = f'{cot}{hang_tong}'
        ct = cong_thuc.get(o)
        if not ct:
            loi.append(f'{o}: không có công thức nào')
            continue
        m = MAU_COUNTA.match(ct.strip())
        if not m:
            loi.append(f'{o}: không phải COUNTA một vùng — {ct!r}')
            continue
        c1, r1, c2, r2 = m.group(1).upper(), int(m.group(2)), \
            m.group(3).upper(), int(m.group(4))
        if c1 != cot or c2 != cot:
            loi.append(f'{o}: đếm nhầm cột {c1}:{c2}, phải là {cot}')
        if r1 != hang_dau:
            loi.append(
                f'{o}: bắt đầu từ dòng {r1}, phải là {hang_dau} — '
                f'bỏ sót {max(0, r1 - hang_dau)} dòng đầu bảng')
        if r2 < hang_cuoi_ghi:
            loi.append(
                f'{o}: dừng ở dòng {r2}, chưa phủ hết dòng {hang_cuoi_ghi}')
        ket.append(r2)
    if ket and len(set(ket)) > 1:
        loi.append(f'Sáu ô dừng ở các dòng khác nhau: {sorted(set(ket))}')
    return loi


def vá_một_cột(nguon, dich, o, cong_thuc_moi):
    """Chép tệp rồi thay ĐÚNG một công thức, giữ nguyên mọi byte khác.

    Mở như tệp zip và vá thẳng XML — KHÔNG dùng openpyxl ghi đè cả workbook
    (CLAUDE.md đã ghi vì sao). Đây cũng đúng cách một lỗi hồi quy thật sẽ len
    vào: ai đó sửa mẫu bằng Excel rồi lưu lại.
    """
    with zipfile.ZipFile(nguon) as z:
        muc = z.namelist()
        noi_dung = {n: z.read(n) for n in muc}
    x = noi_dung[SHEET_XML].decode('utf-8')
    cu = re.search(r'(<c r="%s"[^>]*>\s*<f>)([^<]*)(</f>)' % re.escape(o), x)
    if not cu:
        return False
    x = x[:cu.start(2)] + cong_thuc_moi + x[cu.end(2):]
    noi_dung[SHEET_XML] = x.encode('utf-8')
    with zipfile.ZipFile(dich, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in muc:
            z.writestr(n, noi_dung[n])
    return True


def bo_cuc_tu_cau_hinh():
    """Suy dòng đầu / dòng cuối / các cột TỪ cauhinh.json, không ghim số trần."""
    cfg = json.load(io.open(os.path.join(GOC, 'cauhinh.json'), encoding='utf-8'))
    tt = cfg.get('thong_tin') or {}
    vt = tt.get('vi_tri_bang') or {}
    kg = cfg.get('khung_gio') or {}
    hang_dau = vt.get('hang_bat_dau')
    cot_dau = vt.get('cot_bat_dau')
    so_lop = len(cfg.get('lop') or [])
    so_ngay = len(kg.get('cac_thu') or [])
    so_tiet = len(kg.get('cac_tiet') or [])
    return cfg, hang_dau, cot_dau, so_lop, so_ngay, so_tiet


def ten_cot(i):
    """1 -> A. Đủ dùng cho bảng vài chục cột."""
    s = ''
    while i > 0:
        i, d = divmod(i - 1, 26)
        s = chr(65 + d) + s
    return s


# ==========================================================================
def ca_0_bo_do_biet_bat_va_biet_bo_qua():
    print('\nCA 0 — bộ dò biết BẮT và biết BỎ QUA (hai chiều)')
    print('-' * 74)
    COT = ['D', 'E']
    # chiều BỎ QUA: vùng đúng thì không phàn nàn gì
    sach = {'D43': 'COUNTA(D8:D42)', 'E43': 'COUNTA($E$8:$E$42)'}
    kiem(kiem_vung(sach, COT, 43, 8, 32) == [],
         'Vùng đúng (kể cả dạng $D$8) -> bộ dò im lặng')
    # chiều BẮT: đúng hình dạng lỗi của trường
    hong = {'D43': 'COUNTA(D10:D42)', 'E43': 'COUNTA(E11:E42)'}
    ra = kiem_vung(hong, COT, 43, 8, 32)
    kiem(len(ra) >= 2, f'Lỗi của trường (D10, E11) -> bắt được {len(ra)} chỗ')
    kiem(any('bỏ sót 2 dòng' in x for x in ra),
         'và nói đúng bỏ sót MẤY dòng — D10 thiếu 2')
    kiem(any('bỏ sót 3 dòng' in x for x in ra), 'E11 thiếu 3')
    # các kiểu hỏng khác
    kiem(kiem_vung({'D43': 'COUNTA(D8:D42)'}, COT, 43, 8, 32),
         'Thiếu hẳn một ô -> bắt được')
    kiem(kiem_vung({'D43': 'COUNT(D8:D42)', 'E43': 'COUNTA(E8:E42)'},
                   COT, 43, 8, 32),
         'COUNT (không phải COUNTA) -> bắt được')
    kiem(kiem_vung({'D43': 'COUNTA(D8:D32)', 'E43': 'COUNTA(E8:E32)'},
                   COT, 43, 8, 32) == [],
         'Dừng ĐÚNG ở dòng cuối vùng ghi -> chấp nhận (không đòi phải tới 42)')
    kiem(kiem_vung({'D43': 'COUNTA(D8:D32)', 'E43': 'COUNTA(E8:E42)'},
                   COT, 43, 8, 32),
         'Hai ô dừng ở HAI dòng khác nhau -> bắt được (mẫu phải nhất quán)')
    kiem(kiem_vung({'D43': 'COUNTA(D8:D31)', 'E43': 'COUNTA(E8:E42)'},
                   COT, 43, 8, 32),
         'Dừng SỚM hơn dòng cuối vùng ghi -> bắt được')
    kiem(kiem_vung({'D43': 'COUNTA(E8:E42)', 'E43': 'COUNTA(E8:E42)'},
                   COT, 43, 8, 32),
         'Đếm nhầm sang cột khác -> bắt được')
    kiem(doc_cong_thuc(os.path.join(GOC, 'khong-co-tep-nay.xlsx')) == {},
         'Tệp không tồn tại -> trả {} chứ KHÔNG ném (hàm toàn phần)')


# ==========================================================================
def ca_1_mau_that_sach():
    print('\nCA 1 — mau_tkb.xlsx thật: sáu ô đếm từ dòng đầu vùng dữ liệu')
    print('-' * 74)
    cfg, hang_dau, cot_dau, so_lop, so_ngay, so_tiet = bo_cuc_tu_cau_hinh()

    # Điều kiện tiên quyết: dữ liệu phải đủ để phép so có nghĩa (bẫy n_uyen == 0)
    kiem(isinstance(hang_dau, int) and hang_dau > 0,
         f'cauhinh.json khai được dòng đầu bảng — {hang_dau}')
    kiem(isinstance(cot_dau, int) and cot_dau > 0,
         f'và cột đầu bảng — {cot_dau}')
    kiem(so_lop > 1, f'và có {so_lop} lớp (phải > 1 thì phép so mới có nghĩa)')
    kiem(so_ngay > 1 and so_tiet > 1,
         f'và khung giờ {so_ngay} ngày x {so_tiet} tiết')
    if not (hang_dau and cot_dau and so_lop and so_ngay and so_tiet):
        kiem(False, 'Thiếu khai báo bố cục — bỏ phần còn lại của CA 1')
        return None

    cac_cot = [ten_cot(cot_dau + j) for j in range(so_lop)]
    hang_cuoi_ghi = hang_dau + so_ngay * so_tiet - 1
    print(f'       suy từ cauhinh.json: cột {cac_cot[0]}..{cac_cot[-1]}, '
          f'dòng {hang_dau}..{hang_cuoi_ghi}')

    mau = os.path.join(GOC, 'mau_tkb.xlsx')
    ct = doc_cong_thuc(mau)
    kiem(len(ct) > 1, f'Đọc được {len(ct)} công thức trong mẫu')

    # dòng tổng: dòng có COUNTA của cột đầu tiên
    hang_tong = None
    for o, f in ct.items():
        m = re.match(r'^([A-Z]{1,3})(\d+)$', o)
        if m and m.group(1) == cac_cot[0] and MAU_COUNTA.match(f.strip()):
            hang_tong = int(m.group(2))
            break
    kiem(hang_tong is not None,
         f'Tìm được dòng TỔNG SỐ TIẾT/TUẦN — dòng {hang_tong}')
    if hang_tong is None:
        return None

    loi = kiem_vung(ct, cac_cot, hang_tong, hang_dau, hang_cuoi_ghi)
    for x in loi:
        print(f'       ! {x}')
    kiem(loi == [],
         f'Cả {len(cac_cot)} ô đếm từ dòng {hang_dau}, phủ hết tới dòng '
         f'{hang_cuoi_ghi} — {len(loi)} chỗ sai')
    return (cac_cot, hang_tong, hang_dau, hang_cuoi_ghi)


# ==========================================================================
def ca_2_tep_xuat_ra_giu_nguyen(bo_cuc, tm):
    print('\nCA 2 — HÀNH VI: xuất Excel thật rồi soi lại tệp ra')
    print('-' * 74)
    if not bo_cuc:
        kiem(False, 'CA 1 không dựng được bố cục — bỏ CA 2')
        return
    cac_cot, hang_tong, hang_dau, hang_cuoi_ghi = bo_cuc

    import xep_tkb as X
    try:
        cfg = X.doc_cau_hinh(os.path.join(GOC, 'cauhinh.json'))
        bt = X.BaiToan(cfg)
    except Exception as e:
        kiem(False, f'Dựng bài toán -> {type(e).__name__}: {e}')
        return
    kiem(len(bt.cac_cum) > 1, f'Dựng được bài toán — {len(bt.cac_cum)} cụm')
    for ten in ('SO_NGAY', 'SO_TIET', 'lop', 'cac_cum'):
        if not hasattr(bt, ten):
            kiem(False, f'BaiToan thiếu thuộc tính {ten} — hình dạng đã đổi')
            return

    # Xếp một phương án GIẢ, đủ hợp lệ để lấp kín lưới: mỗi lớp lấp lần lượt.
    # Không gọi bộ giải — CA này canh việc XUẤT, không canh việc GIẢI.
    pa, o_da_dung = {}, {}
    for c in bt.cac_cum:
        dung = o_da_dung.setdefault(c.lop, [])
        dat = False
        for d in range(bt.SO_NGAY):
            for v in range(bt.SO_TIET - c.size + 1):
                nếu_kẹt = any((d, v + k) in dung for k in range(c.size))
                if nếu_kẹt:
                    continue
                for k in range(c.size):
                    dung.append((d, v + k))
                pa[c.ma] = (d, v)
                dat = True
                break
            if dat:
                break
    kiem(len(pa) == len(bt.cac_cum),
         f'Xếp giả lấp được hết {len(pa)}/{len(bt.cac_cum)} cụm')

    xuat = os.path.join(tm, 'ra.xlsx')
    X.xuat_excel(bt, pa, os.path.join(GOC, 'mau_tkb.xlsx'), xuat)
    kiem(os.path.exists(xuat), 'Xuất được tệp Excel')

    ct = doc_cong_thuc(xuat)
    loi = kiem_vung(ct, cac_cot, hang_tong, hang_dau, hang_cuoi_ghi)
    for x in loi:
        print(f'       ! {x}')
    kiem(loi == [],
         f'Sáu công thức QUA ĐƯỢC bước xuất, còn nguyên vùng — {len(loi)} sai')

    # Vùng COUNTA phủ có đúng bằng số ô chương trình vừa ghi không
    import openpyxl
    wb = openpyxl.load_workbook(xuat)
    tt = cfg.get('thong_tin') or {}
    ws = wb[tt.get('ten_sheet')]
    ky_vong = {}
    for c in bt.cac_cum:
        if not c.la_trong:
            ky_vong[c.lop] = ky_vong.get(c.lop, 0) + c.size
    kiem(len(ky_vong) > 1 and min(ky_vong.values()) > 1,
         f'Số tiết kỳ vọng suy TỪ cauhinh.json, khác 0 — {sorted(set(ky_vong.values()))}')

    lech = []
    for j, cot in enumerate(cac_cot):
        if j >= len(bt.lop):
            break
        lop = bt.lop[j]
        dem = sum(1 for r in range(hang_dau, hang_cuoi_ghi + 1)
                  if ws[f'{cot}{r}'].value not in (None, ''))
        if dem != ky_vong.get(lop):
            lech.append(f'{lop} cột {cot}: đếm {dem}, kỳ vọng {ky_vong.get(lop)}')
    for x in lech:
        print(f'       ! {x}')
    kiem(lech == [],
         f'Mỗi cột có đúng số ô có chữ bằng số tiết của lớp đó — {len(lech)} lệch')

    # và nếu dùng vùng SAI của trường thì con số sẽ TỤT — chứng minh lỗi có thật
    if hang_dau + 2 <= hang_cuoi_ghi:
        cot0 = cac_cot[0]
        dung = sum(1 for r in range(hang_dau, hang_cuoi_ghi + 1)
                   if ws[f'{cot0}{r}'].value not in (None, ''))
        sai = sum(1 for r in range(hang_dau + 2, hang_cuoi_ghi + 1)
                  if ws[f'{cot0}{r}'].value not in (None, ''))
        kiem(sai < dung,
             f'Bỏ 2 dòng đầu thì số tiết TỤT thật: {dung} -> {sai} '
             f'(đúng hình dạng lỗi tệp tuần 1 của trường)')


# ==========================================================================
def ca_3_lui_ban_va_thi_phai_do(bo_cuc, tm):
    print('\nCA 3 — LÙI BẢN VÁ: đổi một cột về dòng sau, ca thử phải ĐỎ')
    print('-' * 74)
    if not bo_cuc:
        kiem(False, 'CA 1 không dựng được bố cục — bỏ CA 3')
        return
    cac_cot, hang_tong, hang_dau, hang_cuoi_ghi = bo_cuc
    mau = os.path.join(GOC, 'mau_tkb.xlsx')
    cot0 = cac_cot[0]
    o = f'{cot0}{hang_tong}'
    lui = hang_dau + 2                       # đúng kiểu D10 của tệp trường
    ban_hong = os.path.join(tm, 'mau_hong.xlsx')

    ok = vá_một_cột(mau, ban_hong, o,
                    f'COUNTA({cot0}{lui}:{cot0}{hang_cuoi_ghi + 10})')
    kiem(ok, f'Vá được bản chép: {o} -> COUNTA({cot0}{lui}:...)')
    if not ok:
        return

    kiem(os.path.getsize(ban_hong) > 1, 'Bản hỏng là một tệp thật')
    ct = doc_cong_thuc(ban_hong)
    kiem(len(ct) > 1, f'Bản hỏng vẫn mở được, đọc {len(ct)} công thức')

    loi = kiem_vung(ct, cac_cot, hang_tong, hang_dau, hang_cuoi_ghi)
    kiem(loi != [], f'Bộ dò chuyển sang ĐỎ trên bản hỏng — {len(loi)} chỗ')
    kiem(any(o in x and 'bỏ sót 2 dòng' in x for x in loi),
         f'và chỉ ĐÍCH DANH {o}, nói đúng bỏ sót 2 dòng')

    # đỏ ĐÚNG một cột, năm cột kia phải còn sạch
    con_lai = [x for x in loi if not x.startswith(o)]
    kiem(con_lai == [],
         f'Năm cột không bị vá vẫn SẠCH — {len(con_lai)} lời phàn nàn thừa')

    # bản thật không hề bị đụng
    kiem(kiem_vung(doc_cong_thuc(mau), cac_cot, hang_tong,
                   hang_dau, hang_cuoi_ghi) == [],
         'mau_tkb.xlsx THẬT không bị đụng tới — vẫn sạch')

    # và mọi phần khác của tệp phải giống hệt: chỉ đổi đúng một công thức
    with zipfile.ZipFile(mau) as a, zipfile.ZipFile(ban_hong) as b:
        khac = [n for n in a.namelist()
                if n != SHEET_XML and a.read(n) != b.read(n)]
    kiem(khac == [], f'Ngoài sheet đó, mọi mục khác giống hệt — {khac[:3]}')


# ==========================================================================
def ca_4_a8_khong_bat_duoc_loi_nay():
    print('\nCA 4 — vì sao CA A8 của thu_an_ninh.py KHÔNG thay được tệp này')
    print('-' * 74)
    # A8 hỏi "bảy công thức CÒN hay MẤT". Ta dựng đúng thứ A8 thấy là ĐỦ,
    # rồi cho bộ dò của tệp này xem: nó phải thấy SAI.
    COT = ['D', 'E', 'F']
    nhu_a8_thay = {f'{c}43' for c in COT}
    hong = {'D43': 'COUNTA(D10:D42)', 'E43': 'COUNTA(E11:E42)',
            'F43': 'COUNTA(F10:F42)'}
    kiem(set(hong) == nhu_a8_thay,
         'Bản hỏng vẫn có ĐỦ ô ở dòng 43 — A8 đếm ô thì thấy đủ, vẫn xanh')
    kiem(all(x.upper().startswith('COUNTA') for x in hong.values()),
         'và vẫn là COUNTA — A8 nhìn tên hàm cũng thấy đúng')
    ra = kiem_vung(hong, COT, 43, 8, 32)
    kiem(len(ra) >= 3,
         f'Nhưng bộ dò của tệp này ĐỎ ngay — {len(ra)} chỗ sai')
    print('       => hai ca thử canh hai chuyện khác nhau, cần cả hai.')


# ==========================================================================
def main():
    dat_dong_ho_canh(180)
    print('=' * 74)
    print(' THỬ: sáu ô "TỔNG SỐ TIẾT/TUẦN" đếm từ đúng dòng đầu vùng dữ liệu')
    print('=' * 74)
    tm = tempfile.mkdtemp(prefix='thu_o_tong_')
    try:
        ca_0_bo_do_biet_bat_va_biet_bo_qua()
        bo_cuc = ca_1_mau_that_sach()
        ca_2_tep_xuat_ra_giu_nguyen(bo_cuc, tm)
        ca_3_lui_ban_va_thi_phai_do(bo_cuc, tm)
        ca_4_a8_khong_bat_duoc_loi_nay()
    finally:
        shutil.rmtree(tm, ignore_errors=True)

    dat = sum(1 for ok, _ in KQ if ok)
    print('\n' + '=' * 74)
    print(f'Tổng: {dat}/{len(KQ)} mục ĐẠT.')
    return 0 if dat == len(KQ) else 1


if __name__ == '__main__':
    sys.exit(main())
