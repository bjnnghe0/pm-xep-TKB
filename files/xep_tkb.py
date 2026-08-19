#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
==========================================================================
 PHẦN MỀM XẾP THỜI KHÓA BIỂU  --  Khoa Khoa học Cơ bản
 Trường Cao đẳng Đà Lạt
--------------------------------------------------------------------------
 Cách dùng:
     python xep_tkb.py                       # chạy với cauhinh.json
     python xep_tkb.py --cauhinh khac.json   # dùng file cấu hình khác
     python xep_tkb.py --seed 123 --lan 800  # đổi seed / số lần thử
     python xep_tkb.py --xuat TKB_Tuan3.xlsx # đổi tên file kết quả

 Toàn bộ dữ liệu (lớp, giáo viên, số tiết, ràng buộc) nằm trong
 cauhinh.json - sửa bằng Notepad là đủ, không cần đụng vào file này.
==========================================================================
"""

import json
import time
import random
import argparse
import shutil
import sys
import os
from collections import defaultdict

# ==========================================================================
# 1. ĐỌC CẤU HÌNH
# ==========================================================================

def doc_cau_hinh(duong_dan):
    with open(duong_dan, encoding='utf-8') as f:
        return json.load(f)


# ==========================================================================
# 2. MÔ HÌNH DỮ LIỆU
# ==========================================================================

class Cum:
    """Một 'cụm tiết': nhóm tiết liên tiếp của cùng 1 môn, 1 GV, 1 lớp.
       VD: Vật lý 2 tiết liền -> 1 cụm size=2."""

    __slots__ = ('ma', 'lop', 'mon', 'nhom', 'gv', 'size', 'la_cd', 'la_trong',
                 'ngay_ok', 'vitri_ok', 'ngay', 'vitri', 'luat_khac_ngay')

    def __init__(self, ma, lop, mon, nhom, gv, size,
                 la_cd=False, la_trong=False):
        self.ma = ma
        self.lop = lop
        self.mon = mon
        self.nhom = nhom
        self.gv = gv
        self.size = size
        self.la_cd = la_cd            # là tiết chuyên đề?
        self.la_trong = la_trong      # là ô trống (không có tiết)?
        self.ngay_ok = set()          # các thứ được phép (index 0..n-1)
        self.vitri_ok = set()         # vị trí bắt đầu được phép (0..n-1)
        self.ngay = None
        self.vitri = None
        self.luat_khac_ngay = []      # [(chi_so_luat, 'a'|'b'), ...]

    def nhan(self):
        if self.la_trong:
            return ''
        return f'{self.mon} ({self.gv})'

    def __repr__(self):
        return f'<{self.lop} {self.mon} {self.gv} x{self.size}>'


def chia_cum(so_tiet, mau_cum):
    """Chia số tiết/tuần thành các cụm theo mẫu khai báo trong cauhinh.json.
       VD: so_tiet=3, mau_cum=[2,1] -> [2,1]. Nếu lệch thì tự bù cho đủ."""
    cum, con = [], so_tiet
    for k in mau_cum:
        if con <= 0:
            break
        lay = min(k, con)
        cum.append(lay)
        con -= lay
    while con > 0:            # còn dư thì thêm cụm 1 tiết
        cum.append(1)
        con -= 1
    return cum


class BaiToan:
    """Gom toàn bộ dữ liệu + ràng buộc đã 'dịch' sang dạng máy hiểu được."""

    def __init__(self, cfg):
        self.cfg = cfg
        kg = cfg['khung_gio']
        self.cac_thu = kg['cac_thu']              # [2,3,4,5,6]
        self.cac_tiet = kg['cac_tiet']            # [6,7,8,9,10]
        self.buoi = kg['buoi']
        self.SO_NGAY = len(self.cac_thu)
        self.SO_TIET = len(self.cac_tiet)

        self.lop = [l['ten'] for l in cfg['lop']]
        self.thong_tin_lop = {l['ten']: l for l in cfg['lop']}
        self.dinh_muc = cfg['dinh_muc_mon']
        self.luat = cfg['rang_buoc'].get('luat_chung', {})
        self.uu_tien = cfg.get('uu_tien', {})

        self.cac_cum = []
        self._tao_cum()
        self._ap_rang_buoc()
        self._kiem_tra_so_bo()

    # ---------------- tạo danh sách cụm ----------------
    def _tao_cum(self):
        ma = 0
        for lop in self.lop:
            tong = 0
            for mon, gv, st in self.cfg['phan_cong'][lop]:
                if mon not in self.dinh_muc:
                    raise SystemExit(f'[LỖI] Môn "{mon}" chưa khai báo trong '
                                     f'"dinh_muc_mon" (lớp {lop}).')
                dm = self.dinh_muc[mon]
                for size in chia_cum(st, dm['cum']):
                    c = Cum(ma, lop, mon, dm['nhom'], gv, size,
                            la_cd=dm.get('la_chuyen_de', False))
                    self.cac_cum.append(c)
                    ma += 1
                tong += st

            # ---- ô trống: gom vào cuối 1 buổi để không có tiết trống giữa buổi
            so_o = self.SO_NGAY * self.SO_TIET
            trong = so_o - tong
            if trong < 0:
                raise SystemExit(
                    f'[LỖI] Lớp {lop} có {tong} tiết/tuần nhưng TKB chỉ có '
                    f'{so_o} ô. Giảm số tiết hoặc mở thêm buổi trong "khung_gio".')
            if trong > 0:
                c = Cum(ma, lop, '', 'TRỐNG', '', trong, la_trong=True)
                self.cac_cum.append(c)
                ma += 1

    # ---------------- dịch ràng buộc ----------------
    def _ap_rang_buoc(self):
        rb = self.cfg['rang_buoc']
        thu2idx = {t: i for i, t in enumerate(self.cac_thu)}
        tiet2idx = {t: i for i, t in enumerate(self.cac_tiet)}

        # mặc định: mọi cụm đều có thể vào mọi ngày / mọi vị trí hợp lệ
        for c in self.cac_cum:
            c.ngay_ok = set(range(self.SO_NGAY))
            c.vitri_ok = set(range(self.SO_TIET - c.size + 1))

        def khop(c, r):
            if 'mon' in r and c.mon != r['mon']:
                return False
            if 'nhom_mon' in r and c.nhom != r['nhom_mon']:
                return False
            if 'gv' in r and c.gv != r['gv']:
                return False
            if 'lop' in r:
                ds = r['lop'] if isinstance(r['lop'], list) else [r['lop']]
                if c.lop not in ds:
                    return False
            return True

        # (a) chỉ xếp vào các thứ ...
        for r in rb.get('chi_xep_thu', []):
            cho_phep = {thu2idx[t] for t in r['thu'] if t in thu2idx}
            for c in self.cac_cum:
                if not c.la_trong and khop(c, r):
                    c.ngay_ok &= cho_phep

        # (b) không xếp vào các thứ ...
        for r in rb.get('khong_xep_thu', []):
            cam = {thu2idx[t] for t in r['thu'] if t in thu2idx}
            for c in self.cac_cum:
                if not c.la_trong and khop(c, r):
                    c.ngay_ok -= cam

        # (c) không xếp vào các tiết ...
        for r in rb.get('khong_xep_tiet', []):
            cam = set()
            for t in r.get('tiet_trong_buoi', []):        # thứ tự trong buổi
                if 1 <= t <= self.SO_TIET:
                    cam.add(t - 1)
            for t in r.get('tiet', []):                   # số hiệu tiết tuyệt đối
                if t in tiet2idx:
                    cam.add(tiet2idx[t])
            for c in self.cac_cum:
                if not c.la_trong and khop(c, r):
                    c.vitri_ok = {v for v in c.vitri_ok
                                  if not (set(range(v, v + c.size)) & cam)}

        # (d) SHCN xếp tiết cuối buổi
        if self.luat.get('shcn_tiet_cuoi', False):
            for c in self.cac_cum:
                if c.nhom == 'SHCN':
                    c.vitri_ok &= {self.SO_TIET - c.size}

        # (e) ô trống luôn nằm cuối buổi (chống tiết trống giữa buổi)
        if self.luat.get('khong_tiet_trong_giua_buoi', True):
            for c in self.cac_cum:
                if c.la_trong:
                    c.vitri_ok = {self.SO_TIET - c.size}

        # (f) ràng buộc "không cùng ngày"
        self.luat_khac_ngay = rb.get('khac_ngay', [])
        for i, r in enumerate(self.luat_khac_ngay):
            for ben in ('a', 'b'):
                dk = r[ben]
                for c in self.cac_cum:
                    if not c.la_trong and khop(c, dk):
                        c.luat_khac_ngay.append((i, ben))

    def _kiem_tra_so_bo(self):
        loi = [c for c in self.cac_cum if not c.ngay_ok or not c.vitri_ok]
        if loi:
            print('[LỖI] Các cụm sau bị ràng buộc chặn hết mọi vị trí:')
            for c in loi:
                print('   -', c, 'ngày:', sorted(c.ngay_ok),
                      'vị trí:', sorted(c.vitri_ok))
            raise SystemExit('Hãy nới lỏng ràng buộc trong cauhinh.json.')


# ==========================================================================
# 3. BỘ GIẢI (quay lui + heuristic + khởi động lại ngẫu nhiên)
# ==========================================================================

class BoGiai:

    def __init__(self, bt: BaiToan, rnd: random.Random,
                 gioi_han_buoc=30000, gioi_han_giay=3.0):
        self.bt = bt
        self.rnd = rnd
        self.gioi_han = gioi_han_buoc
        self.gioi_han_giay = gioi_han_giay

    # ------- khởi tạo trạng thái rỗng -------
    def _khoi_tao(self):
        bt = self.bt
        self.lich = {l: [[None] * bt.SO_TIET for _ in range(bt.SO_NGAY)]
                     for l in bt.lop}
        self.gv_ban = defaultdict(lambda: [[None] * bt.SO_TIET
                                           for _ in range(bt.SO_NGAY)])
        self.mon_ngay = defaultdict(int)      # (lop, mon, ngay) -> số cụm
        self.ngay_luat = defaultdict(int)     # (chi_so_luat, ben, ngay) -> đếm
        self.buoc = 0

    # ------- kiểm tra đặt được không -------
    def dat_duoc(self, c: Cum, ngay, vitri):
        bt = self.bt
        if ngay not in c.ngay_ok or vitri not in c.vitri_ok:
            return False
        hang_lop = self.lich[c.lop][ngay]
        for p in range(vitri, vitri + c.size):
            if hang_lop[p] is not None:
                return False
        if not c.la_trong:
            hang_gv = self.gv_ban[c.gv][ngay]
            for p in range(vitri, vitri + c.size):
                if hang_gv[p] is not None:
                    return False
            if bt.luat.get('mot_cum_moi_mon_moi_ngay', True):
                if self.mon_ngay[(c.lop, c.mon, ngay)] >= 1:
                    return False
            for (i, ben) in c.luat_khac_ngay:
                doi = 'b' if ben == 'a' else 'a'
                if self.ngay_luat[(i, doi, ngay)] > 0:
                    return False
        return True

    def dat(self, c, ngay, vitri):
        c.ngay, c.vitri = ngay, vitri
        for p in range(vitri, vitri + c.size):
            self.lich[c.lop][ngay][p] = c
            if not c.la_trong:
                self.gv_ban[c.gv][ngay][p] = c
        if not c.la_trong:
            self.mon_ngay[(c.lop, c.mon, ngay)] += 1
            for (i, ben) in c.luat_khac_ngay:
                self.ngay_luat[(i, ben, ngay)] += 1

    def go(self, c):
        ngay, vitri = c.ngay, c.vitri
        for p in range(vitri, vitri + c.size):
            self.lich[c.lop][ngay][p] = None
            if not c.la_trong:
                self.gv_ban[c.gv][ngay][p] = None
        if not c.la_trong:
            self.mon_ngay[(c.lop, c.mon, ngay)] -= 1
            for (i, ben) in c.luat_khac_ngay:
                self.ngay_luat[(i, ben, ngay)] -= 1
        c.ngay = c.vitri = None

    def lua_chon(self, c):
        return [(d, v) for d in c.ngay_ok for v in c.vitri_ok
                if self.dat_duoc(c, d, v)]

    def _ke_tiet_chinh(self, c, ngay, vitri):
        """Cụm chuyên đề có nằm sát ngay trước/sau tiết chính cùng nhóm môn?"""
        hang = self.lich[c.lop][ngay]
        for p in (vitri - 1, vitri + c.size):
            if 0 <= p < self.bt.SO_TIET:
                k = hang[p]
                if k is not None and not k.la_trong and not k.la_cd \
                        and k.nhom == c.nhom:
                    return True
        return False

    def _sap_ung_vien(self, c, lc):
        """Thứ tự thử: chuyên đề ưu tiên ô nằm sát tiết chính (tiêu chí mềm)."""
        if not c.la_cd:
            self.rnd.shuffle(lc)
            return lc
        ke = [x for x in lc if self._ke_tiet_chinh(c, x[0], x[1])]
        xa = [x for x in lc if x not in ke]
        self.rnd.shuffle(ke)
        self.rnd.shuffle(xa)
        return ke + xa

    # ------- quay lui với MRV + forward-checking -------
    def _quay_lui(self, con_lai):
        if not con_lai:
            return True
        self.buoc += 1
        if self.buoc > self.gioi_han:
            raise TimeoutError
        if self.buoc % 512 == 0 and time.time() - self.moc_gio > self.gioi_han_giay:
            raise TimeoutError

        # MRV: chọn cụm ít lựa chọn nhất, ưu tiên cụm dài (khó xếp hơn)
        tot, tot_lc = None, None
        for c in con_lai:
            lc = self.lua_chon(c)
            if not lc:
                return False                     # forward-checking
            if tot is None or (len(lc), c.la_cd, -c.size) < (
                    len(tot_lc), tot.la_cd, -tot.size):
                tot, tot_lc = c, lc
                if len(lc) == 1:
                    break

        con_lai.remove(tot)
        for (d, v) in self._sap_ung_vien(tot, tot_lc):
            if self.dat_duoc(tot, d, v):
                self.dat(tot, d, v)
                if self._quay_lui(con_lai):
                    return True
                self.go(tot)
        con_lai.append(tot)
        return False

    def giai(self):
        self.moc_gio = time.time()
        self._khoi_tao()
        for c in self.bt.cac_cum:
            c.ngay = c.vitri = None
        con_lai = list(self.bt.cac_cum)
        # xếp cụm dài / bị ràng buộc nhiều trước cho nhanh
        con_lai.sort(key=lambda c: (len(c.ngay_ok) * len(c.vitri_ok), -c.size))
        try:
            ok = self._quay_lui(con_lai)
        except TimeoutError:
            return None
        if not ok:
            return None
        return {c.ma: (c.ngay, c.vitri) for c in self.bt.cac_cum}


# ==========================================================================
# 4. CHẤM ĐIỂM PHƯƠNG ÁN (tiêu chí mềm)
# ==========================================================================

def cham_diem(bt: BaiToan, pa, chi_tiet=False):
    ut = bt.uu_tien
    diem = 0
    ct = {}
    theo_ma = {c.ma: c for c in bt.cac_cum}
    for ma, (d, v) in pa.items():
        theo_ma[ma].ngay, theo_ma[ma].vitri = d, v

    # (1) chuyên đề nằm kề tiết chính cùng nhóm môn
    thuong = ut.get('chuyen_de_ke_tiet_chinh', 0)
    if thuong:
        dem_ke, tong_cd = 0, 0
        for c in bt.cac_cum:
            if not c.la_cd:
                continue
            tong_cd += 1
            for k in bt.cac_cum:
                if (k is not c and k.lop == c.lop and k.nhom == c.nhom
                        and not k.la_cd and k.ngay == c.ngay):
                    if (k.vitri + k.size == c.vitri) or (c.vitri + c.size == k.vitri):
                        diem += thuong
                        dem_ke += 1
                        break
        ct['Chuyên đề xếp liền tiết chính'] = f'{dem_ke}/{tong_cd} tiết'

    # (2) giáo viên đến trường ít ngày nhất có thể
    phat = ut.get('giam_so_ngay_den_truong_cua_gv', 0)
    if phat:
        ngay_gv = defaultdict(set)
        for c in bt.cac_cum:
            if not c.la_trong:
                ngay_gv[c.gv].add(c.ngay)
        tong_ngay = sum(len(s) for s in ngay_gv.values())
        diem -= phat * tong_ngay
        ct['Tổng lượt (giáo viên × ngày) đến trường'] = (
            f'{tong_ngay} lượt / {len(ngay_gv)} GV '
            f'(TB {tong_ngay / max(1, len(ngay_gv)):.1f} ngày/GV)')

    # (3) các cụm cùng môn của 1 lớp nên cách nhau >= 2 ngày
    thuong = ut.get('trai_deu_mon_trong_tuan', 0)
    if thuong:
        nhom = defaultdict(list)
        for c in bt.cac_cum:
            if not c.la_trong and not c.la_cd:
                nhom[(c.lop, c.mon)].append(c.ngay)
        dem_deu, tong_cap = 0, 0
        for ds in nhom.values():
            ds.sort()
            for i in range(len(ds) - 1):
                tong_cap += 1
                if ds[i + 1] - ds[i] >= 2:
                    diem += thuong
                    dem_deu += 1
        ct['Các tiết cùng môn cách nhau ≥ 2 ngày'] = f'{dem_deu}/{tong_cap} cặp'
    return (diem, ct) if chi_tiet else diem


# ==========================================================================
# 5. HẬU KIỂM - đối chiếu lại từng ràng buộc trên kết quả cuối
# ==========================================================================

def hau_kiem(bt: BaiToan, pa):
    """Trả về (danh_sach_dong_bao_cao, so_loi)."""
    theo_ma = {c.ma: c for c in bt.cac_cum}
    for ma, (d, v) in pa.items():
        theo_ma[ma].ngay, theo_ma[ma].vitri = d, v
    bc, loi = [], 0

    def ghi(ten, ok, chi_tiet=''):
        nonlocal loi
        if not ok:
            loi += 1
        bc.append(('ĐẠT' if ok else 'LỖI', ten, chi_tiet))

    thu = bt.cac_thu
    ten_thu = lambda i: f'Thứ {thu[i]}'
    ten_tiet = lambda v: f'tiết {bt.cac_tiet[v]}'

    # 1. trùng lịch giáo viên
    ban = defaultdict(list)
    for c in bt.cac_cum:
        if c.la_trong:
            continue
        for p in range(c.vitri, c.vitri + c.size):
            ban[(c.gv, c.ngay, p)].append(c)
    trung = [k for k, v in ban.items() if len(v) > 1]
    ghi('Không giáo viên nào bị trùng 2 lớp cùng lúc', not trung,
        '; '.join(f'{k[0]} {ten_thu(k[1])} {ten_tiet(k[2])}' for k in trung))

    # 2. đủ số tiết theo định mức
    dem = defaultdict(int)
    for c in bt.cac_cum:
        if not c.la_trong:
            dem[(c.lop, c.mon)] += c.size
    thieu = []
    for lop, ds in bt.cfg['phan_cong'].items():
        if lop.startswith('_'):
            continue
        for mon, gv, st in ds:
            if dem[(lop, mon)] != st:
                thieu.append(f'{lop}/{mon}: {dem[(lop, mon)]}≠{st}')
    ghi('Đủ định mức tiết cho mỗi lớp / mỗi môn', not thieu, '; '.join(thieu))

    # 3. không có tiết trống giữa buổi
    ho = []
    for lop in bt.lop:
        for d in range(bt.SO_NGAY):
            hang = [None] * bt.SO_TIET
            for c in bt.cac_cum:
                if c.lop == lop and c.ngay == d and not c.la_trong:
                    for p in range(c.vitri, c.vitri + c.size):
                        hang[p] = c
            day = [i for i, x in enumerate(hang) if x is not None]
            if day and day != list(range(len(day))):
                ho.append(f'{lop} {ten_thu(d)}')
    ghi('Không có tiết trống chen giữa buổi', not ho, '; '.join(ho))

    # 4. quy tắc gom tiết (Lý / Sinh / HĐTN)
    for ten_nhom, so in (('Lý', 2), ('Sinh', 2), ('HĐTN', 3)):
        xau = []
        for lop in bt.lop:
            ds = [c for c in bt.cac_cum if c.lop == lop and c.nhom == ten_nhom]
            if not ds:
                continue
            if len(ds) > 1 or ds[0].size != so:
                xau.append(lop)
        ghi(f'{ten_nhom}: gom đủ {so} tiết trong cùng 1 buổi', not xau,
            '; '.join(xau))

    # 5. các ràng buộc khai báo trong cauhinh.json
    rb = bt.cfg['rang_buoc']
    thu2idx = {t: i for i, t in enumerate(bt.cac_thu)}

    def khop(c, r):
        if 'mon' in r and c.mon != r['mon']:
            return False
        if 'nhom_mon' in r and c.nhom != r['nhom_mon']:
            return False
        if 'gv' in r and c.gv != r['gv']:
            return False
        if 'lop' in r:
            ds = r['lop'] if isinstance(r['lop'], list) else [r['lop']]
            if c.lop not in ds:
                return False
        return True

    for r in rb.get('chi_xep_thu', []):
        ok_set = {thu2idx[t] for t in r['thu'] if t in thu2idx}
        sai = [f'{c.lop} {ten_thu(c.ngay)}' for c in bt.cac_cum
               if not c.la_trong and khop(c, r) and c.ngay not in ok_set]
        ghi(r.get('_ghichu', 'chỉ xếp thứ'), not sai, '; '.join(sai))

    for r in rb.get('khong_xep_thu', []):
        cam = {thu2idx[t] for t in r['thu'] if t in thu2idx}
        sai = [f'{c.lop} {c.mon} {ten_thu(c.ngay)}' for c in bt.cac_cum
               if not c.la_trong and khop(c, r) and c.ngay in cam]
        ghi(r.get('_ghichu', 'không xếp thứ'), not sai, '; '.join(sai))

    for r in rb.get('khong_xep_tiet', []):
        cam = {t - 1 for t in r.get('tiet_trong_buoi', [])}
        cam |= {bt.cac_tiet.index(t) for t in r.get('tiet', []) if t in bt.cac_tiet}
        sai = []
        for c in bt.cac_cum:
            if not c.la_trong and khop(c, r):
                if set(range(c.vitri, c.vitri + c.size)) & cam:
                    sai.append(f'{c.lop} {c.mon} {ten_thu(c.ngay)} {ten_tiet(c.vitri)}')
        ghi(r.get('_ghichu', 'không xếp tiết'), not sai, '; '.join(sai))

    for i, r in enumerate(rb.get('khac_ngay', [])):
        na = {c.ngay for c in bt.cac_cum if not c.la_trong and khop(c, r['a'])}
        nb = {c.ngay for c in bt.cac_cum if not c.la_trong and khop(c, r['b'])}
        chung = na & nb
        ghi(r.get('_ghichu', 'không cùng ngày'), not chung,
            '; '.join(ten_thu(d) for d in sorted(chung)))

    # 6. SHCN tiết cuối
    if bt.luat.get('shcn_tiet_cuoi'):
        sai = [c.lop for c in bt.cac_cum
               if c.nhom == 'SHCN' and c.vitri != bt.SO_TIET - 1]
        ghi('SHCN luôn ở tiết cuối buổi', not sai, '; '.join(sai))

    return bc, loi


# ==========================================================================
# 6. XUẤT EXCEL (ghi đè lên file mẫu, giữ nguyên 100% định dạng)
# ==========================================================================

def xuat_excel(bt: BaiToan, pa, file_mau, file_xuat):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, Border, Side
    except ImportError:
        raise SystemExit('[LỖI] Thiếu thư viện openpyxl. Chạy: pip install openpyxl')

    if not os.path.exists(file_mau):
        raise SystemExit(f'[LỖI] Không tìm thấy file mẫu: {file_mau}')
    shutil.copy(file_mau, file_xuat)

    tt = bt.cfg['thong_tin']
    vt = tt['vi_tri_bang']
    wb = openpyxl.load_workbook(file_xuat)
    ws = wb[tt['ten_sheet']]

    ws.cell(vt['hang_tieu_de'], 1).value = tt['tieu_de']
    ws.cell(vt['hang_tuan'], 1).value = tt['tuan']
    for j, lop in enumerate(bt.lop):
        ws.cell(vt['hang_gvcn'], vt['cot_bat_dau'] + j).value = \
            bt.thong_tin_lop[lop]['gvcn']

    theo_ma = {c.ma: c for c in bt.cac_cum}
    for ma, (d, v) in pa.items():
        theo_ma[ma].ngay, theo_ma[ma].vitri = d, v

    for c in bt.cac_cum:
        if c.la_trong:
            continue
        cot = vt['cot_bat_dau'] + bt.lop.index(c.lop)
        for p in range(c.vitri, c.vitri + c.size):
            hang = vt['hang_bat_dau'] + c.ngay * bt.SO_TIET + p
            ws.cell(hang, cot).value = c.nhan()

    # ---------- sheet phụ: TKB theo giáo viên ----------
    _sheet_giao_vien(wb, bt, Font, Alignment, Border, Side)
    # ---------- sheet phụ: kết quả hậu kiểm ----------
    _sheet_kiem_tra(wb, bt, pa, Font, Alignment, Border, Side)

    wb.save(file_xuat)
    return file_xuat


def _khung(Border, Side):
    m = Side(style='thin')
    return Border(left=m, right=m, top=m, bottom=m)


def _sheet_giao_vien(wb, bt, Font, Alignment, Border, Side):
    ten = 'TKB giáo viên'
    if ten in wb.sheetnames:
        del wb[ten]
    ws = wb.create_sheet(ten)
    kh = _khung(Border, Side)
    tit = Font(name='Times New Roman', size=11, bold=True)
    thg = Font(name='Times New Roman', size=10)
    giua = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.cell(1, 1).value = 'THỜI KHÓA BIỂU THEO GIÁO VIÊN — ' + bt.cfg['thong_tin']['tuan']
    ws.cell(1, 1).font = Font(name='Times New Roman', size=13, bold=True)

    ws.cell(3, 1, 'Giáo viên').font = tit
    ws.cell(3, 2, 'Số tiết').font = tit
    ws.cell(3, 1).border = ws.cell(3, 2).border = kh
    ws.cell(3, 1).alignment = ws.cell(3, 2).alignment = giua
    cot = 3
    for d in range(bt.SO_NGAY):
        for p in range(bt.SO_TIET):
            o = ws.cell(3, cot)
            o.value = f'T{bt.cac_thu[d]}\nt{bt.cac_tiet[p]}'
            o.font = tit
            o.alignment = giua
            o.border = kh
            ws.column_dimensions[openpyxl_letter(cot)].width = 11
            cot += 1
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 7

    gvs = sorted({c.gv for c in bt.cac_cum if not c.la_trong})
    hang = 4
    for gv in gvs:
        ds = [c for c in bt.cac_cum if c.gv == gv]
        o = ws.cell(hang, 1, gv); o.font = thg; o.border = kh; o.alignment = giua
        o = ws.cell(hang, 2, sum(c.size for c in ds))
        o.font = tit; o.border = kh; o.alignment = giua
        cot = 3
        luoi = {}
        for c in ds:
            for p in range(c.vitri, c.vitri + c.size):
                luoi[(c.ngay, p)] = f'{c.lop}\n{c.mon}'
        for d in range(bt.SO_NGAY):
            for p in range(bt.SO_TIET):
                o = ws.cell(hang, cot)
                o.value = luoi.get((d, p), '')
                o.font = thg
                o.alignment = giua
                o.border = kh
                cot += 1
        ws.row_dimensions[hang].height = 26
        hang += 1
    ws.freeze_panes = 'C4'


def _sheet_kiem_tra(wb, bt, pa, Font, Alignment, Border, Side):
    from openpyxl.styles import PatternFill
    ten = 'Kiểm tra ràng buộc'
    if ten in wb.sheetnames:
        del wb[ten]
    ws = wb.create_sheet(ten)
    kh = _khung(Border, Side)
    tit = Font(name='Times New Roman', size=11, bold=True)
    thg = Font(name='Times New Roman', size=11)
    xanh = PatternFill('solid', fgColor='C6EFCE')
    do = PatternFill('solid', fgColor='FFC7CE')

    ws.cell(1, 1, 'BẢNG ĐỐI CHIẾU RÀNG BUỘC').font = Font(
        name='Times New Roman', size=13, bold=True)
    for j, h in enumerate(['Kết quả', 'Nội dung ràng buộc', 'Ghi chú'], start=1):
        o = ws.cell(3, j, h); o.font = tit; o.border = kh
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 62
    ws.column_dimensions['C'].width = 55

    bc, _ = hau_kiem(bt, pa)
    hang = 4
    for (kq, ten_rb, ct) in bc:
        o = ws.cell(hang, 1, kq); o.font = tit; o.border = kh
        o.fill = xanh if kq == 'ĐẠT' else do
        o.alignment = Alignment(horizontal='center')
        o = ws.cell(hang, 2, ten_rb); o.font = thg; o.border = kh
        o.alignment = Alignment(wrap_text=True, vertical='top')
        o = ws.cell(hang, 3, ct); o.font = thg; o.border = kh
        o.alignment = Alignment(wrap_text=True, vertical='top')
        hang += 1

    hang += 1
    ws.cell(hang, 1, 'CHẤT LƯỢNG PHƯƠNG ÁN (tiêu chí mềm)').font = Font(
        name='Times New Roman', size=12, bold=True)
    hang += 1
    _, cl = cham_diem(bt, pa, chi_tiet=True)
    for k, v in cl.items():
        o = ws.cell(hang, 2, k); o.font = thg; o.border = kh
        o = ws.cell(hang, 3, v); o.font = tit; o.border = kh
        hang += 1

    hang += 1
    ws.cell(hang, 1, 'SỐ TIẾT / TUẦN CỦA TỪNG GIÁO VIÊN').font = Font(
        name='Times New Roman', size=12, bold=True)
    hang += 1
    tk_gv = defaultdict(lambda: [0, set()])
    for c in bt.cac_cum:
        if not c.la_trong:
            tk_gv[c.gv][0] += c.size
            tk_gv[c.gv][1].add(c.ngay)
    for j, h in enumerate(['Giáo viên', 'Số tiết/tuần', 'Số buổi lên lớp'], start=1):
        o = ws.cell(hang, j, h); o.font = tit; o.border = kh
    hang += 1
    for gv in sorted(tk_gv):
        st, ng = tk_gv[gv]
        for j, val in enumerate([gv, st, len(ng)], start=1):
            o = ws.cell(hang, j, val); o.font = thg; o.border = kh
        hang += 1


def openpyxl_letter(idx):
    from openpyxl.utils import get_column_letter
    return get_column_letter(idx)


# ==========================================================================
# 7. IN RA MÀN HÌNH
# ==========================================================================

def in_tkb(bt, pa):
    theo_ma = {c.ma: c for c in bt.cac_cum}
    for ma, (d, v) in pa.items():
        theo_ma[ma].ngay, theo_ma[ma].vitri = d, v
    rong = 26
    print()
    print('THỜI KHÓA BIỂU — ' + bt.cfg['thong_tin']['tuan'])
    print('Thứ  Tiết ' + ''.join(l.ljust(rong) for l in bt.lop))
    print('-' * (10 + rong * len(bt.lop)))
    for d in range(bt.SO_NGAY):
        for p in range(bt.SO_TIET):
            dong = f'{bt.cac_thu[d] if p == 0 else "":<5}{bt.cac_tiet[p]:<5}'
            for lop in bt.lop:
                o = ''
                for c in bt.cac_cum:
                    if (c.lop == lop and c.ngay == d and not c.la_trong
                            and c.vitri <= p < c.vitri + c.size):
                        o = c.nhan()
                        break
                dong += o.ljust(rong)
            print(dong)
        print('-' * (10 + rong * len(bt.lop)))


# ==========================================================================
# 8. HÀM MAIN
# ==========================================================================

def main():
    ap = argparse.ArgumentParser(description='Xếp thời khóa biểu')
    ap.add_argument('--cauhinh', default='cauhinh.json')
    ap.add_argument('--xuat', default=None, help='tên file Excel kết quả')
    ap.add_argument('--seed', type=int, default=None)
    ap.add_argument('--lan', type=int, default=None, help='số lần thử')
    ap.add_argument('--in-man-hinh', action='store_true')
    args = ap.parse_args()

    cfg = doc_cau_hinh(args.cauhinh)
    thu_muc = os.path.dirname(os.path.abspath(args.cauhinh))
    bt = BaiToan(cfg)

    tk = cfg.get('tim_kiem', {})
    so_lan = args.lan or tk.get('so_lan_thu', 300)
    seed = args.seed if args.seed is not None else tk.get('seed', 0)
    gh = tk.get('gioi_han_buoc', 300000)

    print('=' * 70)
    print(' XẾP THỜI KHÓA BIỂU'.center(70))
    print('=' * 70)
    print(f'Số lớp        : {len(bt.lop)}  ({", ".join(bt.lop)})')
    print(f'Khung giờ     : Thứ {bt.cac_thu} | buổi {bt.buoi} | '
          f'tiết {bt.cac_tiet[0]}–{bt.cac_tiet[-1]}')
    print(f'Tổng số cụm   : {len([c for c in bt.cac_cum if not c.la_trong])} '
          f'cụm tiết (+{len([c for c in bt.cac_cum if c.la_trong])} ô trống)')
    print(f'Đang tìm phương án ({so_lan} lần thử, seed={seed}) ...')

    gh_giay = tk.get('gioi_han_giay_moi_lan', 3.0)
    tong_giay = tk.get('gioi_han_tong_giay', 90)
    t0 = time.time()

    tot_nhat, diem_nhat, so_ok, da_thu = None, None, 0, 0
    for i in range(so_lan):
        if time.time() - t0 > tong_giay:
            break
        da_thu += 1
        rnd = random.Random(seed + i)
        pa = BoGiai(bt, rnd, gh, gh_giay).giai()
        if pa is None:
            continue
        so_ok += 1
        d = cham_diem(bt, pa)
        if diem_nhat is None or d > diem_nhat:
            diem_nhat, tot_nhat = d, dict(pa)
    so_lan = da_thu

    if tot_nhat is None:
        print('\n[KHÔNG XẾP ĐƯỢC] Các ràng buộc đang mâu thuẫn nhau.')
        print('Gợi ý: nới 1 trong các ràng buộc ở mục "rang_buoc" của cauhinh.json,')
        print('       hoặc tăng "so_lan_thu" / "gioi_han_buoc" ở mục "tim_kiem".')
        sys.exit(1)

    print(f'-> Tìm được {so_ok}/{so_lan} phương án hợp lệ trong '
          f'{time.time() - t0:.0f}s. Chọn phương án điểm cao nhất ({diem_nhat}).')
    _, chi_tiet_diem = cham_diem(bt, tot_nhat, chi_tiet=True)
    print('\nCHẤT LƯỢNG PHƯƠNG ÁN ĐƯỢC CHỌN')
    print('-' * 70)
    for k, v in chi_tiet_diem.items():
        print(f'  • {k}: {v}')

    bc, loi = hau_kiem(bt, tot_nhat)
    print('\nĐỐI CHIẾU RÀNG BUỘC')
    print('-' * 70)
    for kq, ten, ct in bc:
        print(f'  [{kq}] {ten}' + (f'  --> {ct}' if ct else ''))
    print('-' * 70)
    print(f'Tổng: {len(bc) - loi}/{len(bc)} ràng buộc ĐẠT.')

    if args.in_man_hinh:
        in_tkb(bt, tot_nhat)

    tt = cfg['thong_tin']
    file_mau = os.path.join(thu_muc, tt['file_mau'])
    file_xuat = os.path.join(thu_muc, args.xuat or tt['file_xuat'])
    xuat_excel(bt, tot_nhat, file_mau, file_xuat)
    print(f'\nĐã xuất file: {file_xuat}')
    print('   • Sheet "Khối 12"           : TKB chính (đúng mẫu tuần 1)')
    print('   • Sheet "TKB giáo viên"     : lịch dạy từng thầy cô')
    print('   • Sheet "Kiểm tra ràng buộc": bảng đối chiếu ĐẠT/LỖI')


if __name__ == '__main__':
    main()
